"""PayFlow client onboarding service."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session, joinedload

from app.core.exceptions import ConflictError, NotFoundError, ValidationAppError
from app.infrastructure.database.models import (
    PayflowAccountModel,
    PayflowAiSuggestJobModel,
    PayflowClientFieldMappingModel,
    PayflowClientModel,
    PayflowCommunicationModel,
    PayflowHumanReviewAuditModel,
    PayflowHumanReviewModel,
    PayflowImportErrorModel,
    PayflowImportRunModel,
    PayflowNotificationModel,
    PayflowPortfolioModel,
    PayflowRoleModel,
    PayflowRolePermissionModel,
    PayflowRuleModel,
    PayflowStrategyModel,
    PayflowUserClientAssignmentModel,
    PayflowUserClientPermissionModel,
    PayflowUserMembershipModel,
    UserModel,
)
from app.shared.deletion import record_deletion, snapshot_model
from app.modules.payflow.crm_catalog import (
    ACCOUNT_IMPORT_CURRENCIES,
    BULK_UPLOAD_HEADERS,
    CLIENT_IMPORT_MASTER_PROFILE_HEADERS,
    CLIENT_IMPORT_MASTER_REQUIRED,
    CLIENT_IMPORT_NAME_COLUMNS,
    CRM_CLIENT_HIERARCHY_FIELDS,
    CRM_FILE_BEHAVIOUR,
    CRM_INBOUND_FIELDS,
    CRM_OUTBOUND_FIELDS,
    PAYFLOW_TARGET_FIELDS,
)
from app.shared.enums import (
    PayflowAiMode,
    PayflowBusinessDomain,
    PayflowClientAddedThrough,
    PayflowClientStatus,
    PayflowClientType,
    PayflowConnectionStatus,
    PayflowDataSourceType,
    PayflowMappingStatus,
    PayflowOnboardingStepStatus,
    PayflowPortfolioStatus,
    PayflowRoleCode,
    PayflowRoleScope,
    PayflowRuleStatus,
    PayflowRuleType,
)


_CLIENT_TYPE_LABELS = {
    PayflowClientType.FIRST_PARTY.value: "First Party",
    PayflowClientType.THIRD_PARTY.value: "Third Party",
}
_CLIENT_TYPE_FROM_LABEL = {v.lower(): k for k, v in _CLIENT_TYPE_LABELS.items()}
_CLIENT_TYPE_FROM_LABEL.update({k: k for k in _CLIENT_TYPE_LABELS})

_AI_MODE_LABELS = {
    PayflowAiMode.AUTOPILOT.value: "Autopilot",
    PayflowAiMode.SUPERVISED_AI.value: "Supervised AI",
}
_AI_MODE_FROM_LABEL = {v.lower(): k for k, v in _AI_MODE_LABELS.items()}
_AI_MODE_FROM_LABEL.update({k: k for k in _AI_MODE_LABELS})

_DOMAIN_LABELS = {
    PayflowBusinessDomain.COLLECTIONS.value: "Collections",
}
_DOMAIN_FROM_LABEL = {v.lower(): k for k, v in _DOMAIN_LABELS.items()}
_DOMAIN_FROM_LABEL.update({k: k for k in _DOMAIN_LABELS})

_CONNECTION_LABELS = {
    PayflowConnectionStatus.NOT_CONNECTED.value: "Not Connected",
    PayflowConnectionStatus.CONNECTING.value: "Connecting",
    PayflowConnectionStatus.CONNECTED.value: "Connected",
    PayflowConnectionStatus.CONNECTION_FAILED.value: "Connection Failed",
}

_MAPPING_LABELS = {
    PayflowMappingStatus.MAPPED.value: "Mapped",
    PayflowMappingStatus.NEEDS_ATTENTION.value: "Needs Attention",
    PayflowMappingStatus.UNMAPPED.value: "Unmapped",
    PayflowMappingStatus.VALIDATED.value: "Validated",
}
_MAPPING_FROM_LABEL = {v.lower(): k for k, v in _MAPPING_LABELS.items()}
_MAPPING_FROM_LABEL.update({k: k for k in _MAPPING_LABELS})

_STATUS_LABELS = {
    PayflowClientStatus.DRAFT.value: "Draft",
    PayflowClientStatus.ACTIVE.value: "Active",
}

_ADDED_THROUGH_LABELS = {
    PayflowClientAddedThrough.ADD_CLIENT.value: "Created",
    PayflowClientAddedThrough.FILE_UPLOAD.value: "Created from file upload",
}


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _normalize_code(code: str) -> str:
    return re.sub(r"\s+", "", (code or "").strip())


def _parse_client_type(value: str | None, default: str = PayflowClientType.THIRD_PARTY.value) -> str:
    if not value or not str(value).strip():
        return default
    key = str(value).strip().lower().replace("-", "_").replace(" ", "_")
    # first party / first_party / First Party
    key2 = str(value).strip().lower()
    if key2 in _CLIENT_TYPE_FROM_LABEL:
        return _CLIENT_TYPE_FROM_LABEL[key2]
    if key in _CLIENT_TYPE_FROM_LABEL:
        return _CLIENT_TYPE_FROM_LABEL[key]
    if key in ("first_party", "firstparty"):
        return PayflowClientType.FIRST_PARTY.value
    if key in ("third_party", "thirdparty"):
        return PayflowClientType.THIRD_PARTY.value
    raise ValidationAppError("Client type must be First Party or Third Party")


def _parse_client_type_soft(value: str | None) -> str:
    """Map CRM labels (Lender, Non-Commercial, …) without failing import."""
    try:
        return _parse_client_type(value)
    except ValidationAppError:
        raw = (value or "").strip().lower()
        if "first" in raw:
            return PayflowClientType.FIRST_PARTY.value
        return PayflowClientType.THIRD_PARTY.value


_TRUE_BOOLS = frozenset({"true", "1", "yes", "y"})
_FALSE_BOOLS = frozenset({"false", "0", "no", "n"})
_CLIENT_IMPORT_DATE_FIELDS = (
    "start_date",
    "last_placement",
    "created_at",
    "updated_at",
)


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in _TRUE_BOOLS


def _parse_import_bool(value: Any, *, field: str) -> bool:
    """Strict True/False for import flags (rejects garbage values)."""
    if isinstance(value, bool):
        return value
    text = str(value if value is not None else "").strip().lower()
    if text in _TRUE_BOOLS:
        return True
    if text in _FALSE_BOOLS:
        return False
    if not text:
        raise ValidationAppError(f"{field} is required. Use True or False.")
    raise ValidationAppError(
        f'{field} must be True or False. "{_cell_to_str(value)}" is not valid.'
    )


def _normalize_import_currency(raw: str) -> str:
    """Map CRM currency labels to CAD/USD; reject unsupported values."""
    text = (raw or "").strip()
    if not text:
        raise ValidationAppError("Currency is required")
    upper = text.upper()
    if upper in ACCOUNT_IMPORT_CURRENCIES:
        return upper
    if upper in {"CDN", "C$", "CA$"} or "CANADIAN" in upper or upper == "CAD":
        return "CAD"
    if upper in {"US", "US$", "$"} or "US DOLLAR" in upper or "USD" in upper:
        return "USD"
    raise ValidationAppError(
        f'Currency "{text}" is not supported. Use CAD or USD.'
    )


def _validate_optional_import_date(raw: str, field: str) -> None:
    """Reject non-empty date cells that are not parseable."""
    text = (raw or "").strip()
    if not text:
        return
    if re.fullmatch(r"\d+(\.\d+)?", text):
        # Excel serial date as plain number string — accept as numeric date token
        try:
            serial = float(text)
            if serial < 1 or serial > 60000:
                raise ValidationAppError(
                    f"{field} must be a valid date (YYYY-MM-DD)."
                )
            return
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            datetime.strptime(text[:10], fmt)
            return
        except ValueError:
            continue
    # ISO datetime prefix
    try:
        datetime.fromisoformat(text.replace("Z", "+00:00")[:19])
        return
    except ValueError:
        pass
    raise ValidationAppError(f"{field} must be a valid date (YYYY-MM-DD).")


def _row_first(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        val = (row.get(key) or "").strip()
        if val:
            return val
    return ""


_EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
# Canadian postal: A1A 1A1 (aligned with FE CLIENT_POSTAL_CODE_PATTERN)
_POSTAL_CODE_RE = re.compile(r"^[A-Z]\d[A-Z] \d[A-Z]\d$")


def _normalize_postal_code(raw: str | None) -> str:
    text = (raw or "").strip().upper()
    if not text:
        return ""
    compact = re.sub(r"\s+", "", text)
    if len(compact) == 6:
        return f"{compact[:3]} {compact[3:]}"
    return text


def _is_valid_postal_code(raw: str | None) -> bool:
    cleaned = _normalize_postal_code(raw)
    if not cleaned:
        return True
    return bool(_POSTAL_CODE_RE.match(cleaned))


def _guess_client_import_field(message: str) -> str:
    text = (message or "").lower()
    checks = (
        ("does not belong to master", "client_number"),
        ("was not found", "master_client__client_number"),
        ("cannot have master_client__client_number", "master_client__client_number"),
        ("master clients cannot", "master_client__client_number"),
        ("is_master_client", "is_master_client"),
        ("true or false", "is_master_client"),
        ("master client", "master_client__client_number"),
        ("already belongs to another master", "master_client__client_number"),
        ("already belongs to master", "master_client__client_number"),
        ("already used as a master", "client_number"),
        ("already used as a sub", "client_number"),
        ("already used by a master", "client_number"),
        ("already used by a sub", "client_number"),
        ("cannot be changed", "client_number"),
        ("cannot be removed", "client_number"),
        ("start_date", "start_date"),
        ("last_placement", "last_placement"),
        ("created_at", "created_at"),
        ("updated_at", "updated_at"),
        ("must be a valid date", "start_date"),
        ("currency", "currency__name"),
        ("restricted_access_client", "restricted_access_client"),
        ("sliding_scale", "sliding_scale"),
        ("setup incomplete", "setup"),
        ("branding & channels", "setup"),
        ("supervisor assignment", "setup"),
        ("postal code", "zip_code"),
        ("a1a 1a1", "zip_code"),
        ("client_number", "client_number"),
        ("client_code", "client_code"),
        ("company_name", "company_name"),
        ("short_name", "short_name"),
        ("master_client", "master_client__client_number"),
        ("contact name", "clientdemographiccontactinformation__full_name"),
        ("contact email", "email_address"),
        ("email", "email_address"),
        ("contact phone", "phone_number"),
        ("phone", "phone_number"),
        ("country", "country__name"),
        ("language", "correspondence_language__name"),
        ("client type", "client_type__name"),
        ("client_type", "client_type__name"),
        ("duplicate", "client_code"),
    )
    for needle, field in checks:
        if needle in text:
            return field
    return "File structure"


def _cell_to_str(val: Any) -> str:
    """Normalize Excel/CSV cell values (phones as floats, TRUE/FALSE, etc.)."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return "True" if val else "False"
    if isinstance(val, datetime):
        return val.date().isoformat()
    if hasattr(val, "isoformat") and not isinstance(val, (str, bytes)):
        try:
            return val.isoformat()  # date objects
        except Exception:
            pass
    if isinstance(val, int):
        return str(val)
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        # Scientific / non-integer float — keep compact digit form when phone-like
        text = f"{val:.0f}" if abs(val) >= 1e9 else str(val).strip()
        return text
    text = str(val).strip()
    # Excel often rewrites long phones as 1.4165550101E+10
    if re.fullmatch(r"[+-]?\d+(\.\d+)?[eE][+-]?\d+", text):
        try:
            return str(int(float(text)))
        except ValueError:
            return text
    return text


_PRIMARY_CONTACT_FIELDS = (
    "contact_name",
    "contact_title",
    "contact_email",
    "contact_phone",
    "address_line1",
    "address_line2",
    "city",
    "province_state",
    "country",
    "postal_code",
    "correspondence_language",
    "currency_code",
    "crm_status",
)

_BE_ROOT = Path(__file__).resolve().parents[4]
_CLIENT_UPLOADS = _BE_ROOT / "uploads" / "client-imports"


def _parse_ai_mode(value: str | None, default: str = PayflowAiMode.SUPERVISED_AI.value) -> str:
    if not value or not str(value).strip():
        return default
    key = str(value).strip().lower()
    if key in _AI_MODE_FROM_LABEL:
        return _AI_MODE_FROM_LABEL[key]
    key2 = key.replace(" ", "_").replace("-", "_")
    if key2 in _AI_MODE_FROM_LABEL:
        return _AI_MODE_FROM_LABEL[key2]
    raise ValidationAppError("AI mode must be Autopilot or Supervised AI")


def _parse_domain(value: str | None) -> str:
    if not value or not str(value).strip():
        return PayflowBusinessDomain.COLLECTIONS.value
    key = str(value).strip().lower()
    if key in _DOMAIN_FROM_LABEL:
        return _DOMAIN_FROM_LABEL[key]
    raise ValidationAppError("Business domain must be Collections for Phase 1")


class PayflowClientService:
    def __init__(self, db: Session):
        self.db = db

    # ------------------------------------------------------------------
    # Catalog
    # ------------------------------------------------------------------

    def mapping_catalog(self) -> dict:
        account_fields = [
            {**f, "feed": f.get("feed") or "daily_accounts"} for f in CRM_INBOUND_FIELDS
        ]
        client_fields = list(CRM_CLIENT_HIERARCHY_FIELDS)
        return {
            # Back-compat: `fields` = daily account mappings (existing callers).
            "fields": account_fields,
            "account_fields": account_fields,
            "client_fields": client_fields,
            "outbound_fields": CRM_OUTBOUND_FIELDS,
            "payflow_fields": PAYFLOW_TARGET_FIELDS,
            "governance_rules": self._governance_rule_names(),
            "catalog_version": CRM_FILE_BEHAVIOUR.get("catalog_version", "0.3"),
            "file_behaviour": CRM_FILE_BEHAVIOUR,
            "baseline_status": CRM_FILE_BEHAVIOUR.get("baseline_status"),
            "baseline_note": CRM_FILE_BEHAVIOUR.get("baseline_note"),
        }

    def _governance_rule_names(self) -> list[str]:
        """Active system rules available for client governance selection."""
        rows = (
            self.db.query(PayflowRuleModel.name)
            .filter(
                PayflowRuleModel.status == PayflowRuleStatus.ACTIVE.value,
                PayflowRuleModel.rule_type == PayflowRuleType.SYSTEM.value,
            )
            .order_by(PayflowRuleModel.name.asc())
            .all()
        )
        return [name for (name,) in rows if name]

    def _known_rule_names(self) -> set[str]:
        rows = self.db.query(PayflowRuleModel.name).all()
        return {name for (name,) in rows if name}

    # ------------------------------------------------------------------
    # List / get
    # ------------------------------------------------------------------

    def list_clients(
        self,
        *,
        search: str | None = None,
        status: str | None = None,
        client_type: str | None = None,
        business_domain: str | None = None,
        ai_mode: str | None = None,
        supervisor_user_id: str | None = None,
        client_ids: list[int] | None = None,
    ) -> dict:
        q = self.db.query(PayflowClientModel)
        if client_ids is not None:
            if not client_ids:
                return {"clients": []}
            q = q.filter(PayflowClientModel.id.in_(client_ids))
        if search:
            term = f"%{search.strip()}%"
            q = q.filter(
                (PayflowClientModel.name.ilike(term))
                | (PayflowClientModel.code.ilike(term))
                | (PayflowClientModel.category.ilike(term))
            )
        if status:
            q = q.filter(PayflowClientModel.status == status.strip().lower())
        if client_type:
            q = q.filter(
                PayflowClientModel.client_type == _parse_client_type(client_type)
            )
        if business_domain:
            q = q.filter(
                PayflowClientModel.business_domain == _parse_domain(business_domain)
            )
        if ai_mode:
            q = q.filter(PayflowClientModel.ai_mode == _parse_ai_mode(ai_mode))
        if supervisor_user_id:
            try:
                uid = UUID(supervisor_user_id)
            except ValueError as exc:
                raise ValidationAppError("Invalid supervisor filter") from exc
            q = (
                q.join(PayflowUserClientAssignmentModel)
                .join(PayflowUserMembershipModel)
                .filter(PayflowUserMembershipModel.user_id == uid)
            )

        rows = q.order_by(PayflowClientModel.name).all()
        return {"clients": [self._list_item(c) for c in rows]}

    def get_client(self, client_id: int) -> dict:
        client = self._get_or_404(client_id)
        if self._sync_catalog_mappings(client):
            self.db.commit()
            client = self._get_or_404(client_id)
        return self._detail(client)

    # ------------------------------------------------------------------
    # Create / update
    # ------------------------------------------------------------------

    def create_client(self, payload: dict[str, Any]) -> dict:
        client, _action = self.upsert_client(payload, commit=True)
        return self._detail(client)

    def upsert_client(
        self,
        payload: dict[str, Any],
        *,
        commit: bool = True,
        soft_client_type: bool = False,
    ) -> tuple[PayflowClientModel, str]:
        """Create or update a client.

        Identity is ``crm_client_number`` (client_number). A new number creates a
        new client; an existing number updates other fields. client_number itself
        is locked once set — matching by code must never rewrite it.
        """
        name = (payload.get("name") or "").strip()
        code = _normalize_code(payload.get("code") or "")
        crm_number = (payload.get("crm_client_number") or "").strip() or None
        if not name:
            raise ValidationAppError("Client name is required")
        if not code:
            raise ValidationAppError("Client code / reference is required")

        client = self._resolve_client_for_upsert(crm_number=crm_number, code=code)
        type_parser = _parse_client_type_soft if soft_client_type else _parse_client_type
        ds = self._normalize_data_source(payload.get("data_source_type"), allow_none=True)

        if client is None:
            self._assert_unique_code(code)
            if crm_number:
                self._assert_unique_crm_number(crm_number)
                self._assert_master_number_not_used_as_sub(crm_number)
            added_raw = (payload.get("added_through") or "").strip().lower()
            if added_raw in {
                PayflowClientAddedThrough.ADD_CLIENT.value,
                PayflowClientAddedThrough.FILE_UPLOAD.value,
            }:
                added_through = added_raw
            else:
                added_through = PayflowClientAddedThrough.ADD_CLIENT.value
            client = PayflowClientModel(
                name=name,
                code=code,
                crm_client_number=crm_number,
                category=(payload.get("industry") or payload.get("category") or "").strip()
                or None,
                status=PayflowClientStatus.DRAFT.value,
                client_type=type_parser(payload.get("client_type")),
                business_domain=_parse_domain(payload.get("business_domain")),
                ai_mode=_parse_ai_mode(payload.get("ai_mode")),
                data_source_type=ds or PayflowDataSourceType.FILE.value,
                connection_status=(
                    PayflowConnectionStatus.CONNECTED.value
                    if (ds or PayflowDataSourceType.FILE.value)
                    == PayflowDataSourceType.FILE.value
                    else PayflowConnectionStatus.NOT_CONNECTED.value
                ),
                environment="Sandbox",
                sync_frequency="Every 15 minutes",
                brand_name=(payload.get("brand_name") or "").strip()
                or (payload.get("short_name") or "").strip()
                or "",
                sender_name="",
                email_from="collections@payflow.io",
                sms_sender_id="PAYFLOW",
                channel_email=True,
                channel_sms=True,
                channel_whatsapp=False,
                governance_rules=[],
                added_through=added_through,
                updated_at=_utcnow(),
            )
            self._apply_primary_fields(client, payload, create=True)
            self.db.add(client)
            self.db.flush()
            if ds == PayflowDataSourceType.CRM.value or ds is None:
                self._seed_default_mappings(client)
            action = "created"
        else:
            self._assert_crm_number_immutable(
                existing=client.crm_client_number,
                incoming=crm_number,
                code=code,
                entity="client",
            )
            if not (client.crm_client_number or "").strip() and crm_number:
                self._assert_unique_crm_number(crm_number, exclude_id=client.id)
                self._assert_master_number_not_used_as_sub(crm_number)
                client.crm_client_number = crm_number
            if code.lower() != (client.code or "").lower():
                self._assert_unique_code(code, exclude_id=client.id)
                client.code = code
            client.name = name
            if "industry" in payload or "category" in payload:
                industry = payload.get("industry", payload.get("category"))
                client.category = (str(industry).strip() if industry is not None else None) or None
            if payload.get("client_type") is not None:
                client.client_type = type_parser(payload.get("client_type"))
            if payload.get("business_domain") is not None:
                client.business_domain = _parse_domain(payload.get("business_domain"))
            if payload.get("ai_mode") is not None:
                client.ai_mode = _parse_ai_mode(payload.get("ai_mode"))
            if "data_source_type" in payload:
                self._set_data_source(client, ds)
            if payload.get("brand_name") is not None or payload.get("short_name"):
                client.brand_name = (
                    (payload.get("brand_name") or "").strip()
                    or (payload.get("short_name") or "").strip()
                    or client.brand_name
                )
            self._apply_primary_fields(client, payload, create=False)
            if (
                client.data_source_type == PayflowDataSourceType.CRM.value
                and not (client.field_mappings or [])
            ):
                self._seed_default_mappings(client)
            client.updated_at = _utcnow()
            action = "updated"

        if commit:
            self.db.commit()
            self.db.refresh(client)
        else:
            self.db.flush()
        return client, action

    def update_client(self, client_id: int, payload: dict[str, Any]) -> dict:
        client = self._get_or_404(client_id)
        if client.status == PayflowClientStatus.ACTIVE.value:
            # Allow config updates after activation for non-identity fields;
            # identity uniqueness still enforced.
            pass

        if "crm_client_number" in payload:
            incoming = (
                str(payload["crm_client_number"]).strip()
                if payload.get("crm_client_number")
                else None
            ) or None
            self._assert_crm_number_immutable(
                existing=client.crm_client_number,
                incoming=incoming,
                code=client.code,
                entity="client",
            )
            if not (client.crm_client_number or "").strip() and incoming:
                self._assert_unique_crm_number(incoming, exclude_id=client.id)
                self._assert_master_number_not_used_as_sub(incoming)
                client.crm_client_number = incoming

        if "name" in payload and payload["name"] is not None:
            name = str(payload["name"]).strip()
            if not name:
                raise ValidationAppError("Client name is required")
            client.name = name

        if "code" in payload and payload["code"] is not None:
            code = _normalize_code(str(payload["code"]))
            if not code:
                raise ValidationAppError("Client code / reference is required")
            self._assert_unique_code(code, exclude_id=client.id)
            client.code = code

        if "industry" in payload or "category" in payload:
            industry = payload.get("industry", payload.get("category"))
            client.category = (str(industry).strip() if industry is not None else None) or None

        if "client_type" in payload and payload["client_type"] is not None:
            client.client_type = _parse_client_type(payload["client_type"])

        if "business_domain" in payload and payload["business_domain"] is not None:
            client.business_domain = _parse_domain(payload["business_domain"])

        if "ai_mode" in payload and payload["ai_mode"] is not None:
            client.ai_mode = _parse_ai_mode(payload["ai_mode"])

        if "data_source_type" in payload:
            ds = self._normalize_data_source(payload["data_source_type"], allow_none=True)
            self._set_data_source(client, ds)

        if "connection_status" in payload and payload["connection_status"] is not None:
            if client.data_source_type == PayflowDataSourceType.FILE.value:
                client.connection_status = PayflowConnectionStatus.CONNECTED.value
            else:
                cs = str(payload["connection_status"]).strip().lower().replace(" ", "_")
                label_map = {v.lower().replace(" ", "_"): k for k, v in _CONNECTION_LABELS.items()}
                label_map.update({k: k for k in _CONNECTION_LABELS})
                if cs not in label_map:
                    raise ValidationAppError("Invalid connection status")
                client.connection_status = label_map[cs]

        branding_keys = ("brand_name", "sender_name", "email_from", "sms_sender_id", "channels")
        if any(k in payload for k in branding_keys):
            self._require_draft_for_branding(client)

        for field in (
            "crm_system_name",
            "integration_ref",
            "environment",
            "sync_frequency",
            "brand_name",
            "sender_name",
            "email_from",
            "sms_sender_id",
            *_PRIMARY_CONTACT_FIELDS,
        ):
            if field in payload and payload[field] is not None:
                setattr(client, field, str(payload[field]).strip() or None)

        if "channels" in payload and isinstance(payload["channels"], dict):
            ch = payload["channels"]
            if "email" in ch:
                client.channel_email = bool(ch["email"])
            if "sms" in ch:
                client.channel_sms = bool(ch["sms"])
            if "whatsapp" in ch:
                client.channel_whatsapp = bool(ch["whatsapp"])

        if "governance_rules" in payload and payload["governance_rules"] is not None:
            rules = payload["governance_rules"]
            if not isinstance(rules, list):
                raise ValidationAppError("governance_rules must be a list")
            known = self._known_rule_names()
            # Keep previously saved names even if the rule was later deactivated/renamed away.
            previously = set(client.governance_rules or [])
            unknown = [r for r in rules if r not in known and r not in previously]
            if unknown:
                raise ValidationAppError(f"Unknown governance rule(s): {', '.join(map(str, unknown))}")
            client.governance_rules = list(rules)

        if "mappings" in payload and payload["mappings"] is not None:
            self._replace_mappings(client, payload["mappings"])

        newly_assigned = []
        if "supervisor_user_ids" in payload and payload["supervisor_user_ids"] is not None:
            newly_assigned = self._sync_supervisors(client, payload["supervisor_user_ids"])

        client.updated_at = _utcnow()
        self.db.commit()
        self.db.refresh(client)
        if newly_assigned:
            try:
                from app.modules.payflow.services.notification_service import (
                    PayflowNotificationService,
                )

                svc = PayflowNotificationService(self.db)
                for user in newly_assigned:
                    if user:
                        svc.notify_supervisor_assignment(
                            user=user,
                            client_id=client.id,
                            client_name=client.name,
                            send_mail=True,
                        )
            except Exception:
                pass
        return self._detail(client)

    def update_logo(
        self,
        client_id: int,
        *,
        file_bytes: bytes,
        content_type: str | None,
        filename: str | None,
    ) -> dict:
        from pathlib import Path
        import time

        client = self._get_or_404(client_id)
        self._require_draft_for_branding(client)
        if client.client_type == PayflowClientType.THIRD_PARTY.value:
            raise ValidationAppError(
                "Third Party clients use PayFlow operator branding; client logo upload is not used"
            )
        allowed = {
            "image/jpeg": ".jpg",
            "image/jpg": ".jpg",
            "image/png": ".png",
            "image/webp": ".webp",
            "image/gif": ".gif",
        }
        ext = None
        if content_type and content_type.lower() in allowed:
            ext = allowed[content_type.lower()]
        elif filename:
            lower = filename.lower()
            if lower.endswith(".jpg") or lower.endswith(".jpeg"):
                ext = ".jpg"
            elif lower.endswith(".png"):
                ext = ".png"
            elif lower.endswith(".webp"):
                ext = ".webp"
            elif lower.endswith(".gif"):
                ext = ".gif"
        if not ext:
            raise ValidationAppError("Upload a JPG, PNG, WEBP or GIF image.")
        if len(file_bytes) > 2 * 1024 * 1024:
            raise ValidationAppError("Logo must be 2 MB or smaller.")

        project_root = Path(__file__).resolve().parents[4]
        upload_root = project_root / "uploads" / "client-logos"
        upload_root.mkdir(parents=True, exist_ok=True)

        for old in upload_root.glob(f"{client.id}*"):
            try:
                old.unlink()
            except OSError:
                pass

        dest_name = f"{client.id}-{int(time.time())}{ext}"
        dest = upload_root / dest_name
        dest.write_bytes(file_bytes)
        client.logo_url = f"/uploads/client-logos/{dest_name}"
        client.updated_at = _utcnow()
        self.db.commit()
        self.db.refresh(client)
        return self._detail(client)

    def remove_logo(self, client_id: int) -> dict:
        from pathlib import Path

        client = self._get_or_404(client_id)
        self._require_draft_for_branding(client)
        project_root = Path(__file__).resolve().parents[4]
        upload_root = project_root / "uploads" / "client-logos"
        for old in upload_root.glob(f"{client.id}*"):
            try:
                old.unlink()
            except OSError:
                pass
        client.logo_url = None
        client.updated_at = _utcnow()
        self.db.commit()
        self.db.refresh(client)
        return self._detail(client)

    def activate_client(self, client_id: int, user: UserModel | None = None) -> tuple[dict, list[int]]:
        client = self._get_or_404(client_id)
        blockers = self._activation_blockers(client)
        if blockers:
            raise ValidationAppError("; ".join(blockers))
        client.status = PayflowClientStatus.ACTIVE.value
        client.updated_at = _utcnow()
        self.db.commit()
        self.db.refresh(client)

        from app.modules.payflow.services.ai_suggest_job_service import AiSuggestJobService

        job_ids = AiSuggestJobService(self.db).enqueue_for_client(
            client_id,
            trigger="client_activated",
            requested_by_user_id=user.id if user else None,
        )
        if job_ids:
            self.db.commit()
        return self._detail(client), job_ids

    # ------------------------------------------------------------------
    # Portfolios
    # ------------------------------------------------------------------

    def list_portfolios(self, client_id: int) -> dict:
        self._get_or_404(client_id)
        rows = (
            self.db.query(PayflowPortfolioModel)
            .options(
                joinedload(PayflowPortfolioModel.accounts),
                joinedload(PayflowPortfolioModel.strategies),
            )
            .filter(PayflowPortfolioModel.client_id == client_id)
            .order_by(PayflowPortfolioModel.name)
            .all()
        )
        return {"portfolios": [self._portfolio_item(p) for p in rows]}

    def create_portfolio(self, client_id: int, payload: dict[str, Any]) -> dict:
        row, _action = self.upsert_portfolio(client_id, payload, commit=True)
        return self._portfolio_item(row)

    def upsert_portfolio(
        self,
        client_id: int,
        payload: dict[str, Any],
        *,
        commit: bool = True,
    ) -> tuple[PayflowPortfolioModel, str]:
        client = self._get_or_404(client_id)
        name = (payload.get("name") or "").strip()
        code = _normalize_code(payload.get("code") or "")
        crm_number = (payload.get("crm_client_number") or "").strip() or None
        if not name:
            raise ValidationAppError("Portfolio name is required")
        if not code:
            raise ValidationAppError("Portfolio code / reference is required")

        row = self._resolve_portfolio_for_upsert(
            client.id, crm_number=crm_number, code=code
        )
        status = (payload.get("status") or PayflowPortfolioStatus.ONBOARDING.value).strip().lower()
        if status not in {s.value for s in PayflowPortfolioStatus}:
            raise ValidationAppError("Invalid portfolio status")

        if row is None:
            clash = (
                self.db.query(PayflowPortfolioModel)
                .filter(
                    PayflowPortfolioModel.client_id == client.id,
                    PayflowPortfolioModel.code == code,
                )
                .first()
            )
            if clash:
                raise ConflictError("Portfolio code already exists for this client")
            row = PayflowPortfolioModel(
                client_id=client.id,
                name=name,
                code=code,
                status=status,
                description=(payload.get("description") or "").strip() or None,
                crm_client_number=crm_number,
            )
            self.db.add(row)
            action = "created"
        else:
            self._assert_crm_number_immutable(
                existing=row.crm_client_number,
                incoming=crm_number,
                code=code,
                entity="sub-client",
            )
            if not (row.crm_client_number or "").strip() and crm_number:
                self._assert_sub_import_identity(
                    number=crm_number,
                    master_number=(client.crm_client_number or "").strip()
                    or f"client_id:{client.id}",
                    parent=client,
                )
                row.crm_client_number = crm_number
            if code.lower() != (row.code or "").lower():
                clash = (
                    self.db.query(PayflowPortfolioModel)
                    .filter(
                        PayflowPortfolioModel.client_id == client.id,
                        PayflowPortfolioModel.code == code,
                        PayflowPortfolioModel.id != row.id,
                    )
                    .first()
                )
                if clash:
                    raise ConflictError("Portfolio code already exists for this client")
                row.code = code
            row.name = name
            if "status" in payload and payload["status"] is not None:
                row.status = status
            if "description" in payload:
                row.description = (
                    str(payload["description"]).strip() if payload.get("description") else None
                )
            action = "updated"

        client.updated_at = _utcnow()
        if commit:
            self.db.commit()
            self.db.refresh(row)
        else:
            self.db.flush()
        return row, action

    def update_portfolio(
        self,
        client_id: int,
        portfolio_id: int,
        payload: dict[str, Any],
        user: UserModel | None = None,
    ) -> tuple[dict, list[int]]:
        client = self._get_or_404(client_id)
        row = (
            self.db.query(PayflowPortfolioModel)
            .filter(
                PayflowPortfolioModel.id == portfolio_id,
                PayflowPortfolioModel.client_id == client_id,
            )
            .first()
        )
        if not row:
            raise NotFoundError("Portfolio not found")

        became_active = False
        if "name" in payload and payload["name"] is not None:
            name = str(payload["name"]).strip()
            if not name:
                raise ValidationAppError("Portfolio name is required")
            row.name = name
        if "code" in payload and payload["code"] is not None:
            code = _normalize_code(str(payload["code"]))
            if not code:
                raise ValidationAppError("Portfolio code / reference is required")
            clash = (
                self.db.query(PayflowPortfolioModel)
                .filter(
                    PayflowPortfolioModel.client_id == client_id,
                    PayflowPortfolioModel.code == code,
                    PayflowPortfolioModel.id != portfolio_id,
                )
                .first()
            )
            if clash:
                raise ConflictError("Portfolio code already exists for this client")
            row.code = code
        if "status" in payload and payload["status"] is not None:
            status = str(payload["status"]).strip().lower()
            if status not in {s.value for s in PayflowPortfolioStatus}:
                raise ValidationAppError("Invalid portfolio status")
            prev = (row.status or "").lower()
            row.status = status
            became_active = (
                status == PayflowPortfolioStatus.ACTIVE.value
                and prev != PayflowPortfolioStatus.ACTIVE.value
            )
        if "description" in payload:
            row.description = (
                str(payload["description"]).strip() if payload["description"] else None
            )
        if "crm_client_number" in payload:
            incoming = (
                str(payload["crm_client_number"]).strip()
                if payload.get("crm_client_number")
                else None
            ) or None
            self._assert_crm_number_immutable(
                existing=row.crm_client_number,
                incoming=incoming,
                code=row.code,
                entity="sub-client",
            )
            if not (row.crm_client_number or "").strip() and incoming:
                self._assert_sub_import_identity(
                    number=incoming,
                    master_number=(client.crm_client_number or "").strip()
                    or f"client_id:{client.id}",
                    parent=client,
                )
                row.crm_client_number = incoming

        client.updated_at = _utcnow()
        self.db.commit()
        self.db.refresh(row)

        job_ids: list[int] = []
        if became_active:
            from app.modules.payflow.services.ai_suggest_job_service import AiSuggestJobService

            job = AiSuggestJobService(self.db).enqueue(
                client_id,
                portfolio_id,
                trigger="portfolio_activated",
                requested_by_user_id=user.id if user else None,
            )
            if job:
                job_ids = [job.id]
                self.db.commit()
        return self._portfolio_item(row), job_ids

    def _purge_review(self, review: PayflowHumanReviewModel) -> list[dict[str, str]]:
        related: list[dict[str, str]] = []
        audits = (
            self.db.query(PayflowHumanReviewAuditModel)
            .filter(PayflowHumanReviewAuditModel.review_id == review.id)
            .all()
        )
        for audit in audits:
            self.db.delete(audit)
        related.append({"entity_type": "review", "id": str(review.id), "label": review.code})
        self.db.delete(review)
        return related

    def _purge_ai_suggest_jobs(
        self, *, client_id: int | None = None, portfolio_id: int | None = None
    ) -> list[dict[str, str]]:
        q = self.db.query(PayflowAiSuggestJobModel)
        if client_id is not None:
            q = q.filter(PayflowAiSuggestJobModel.client_id == client_id)
        if portfolio_id is not None:
            q = q.filter(PayflowAiSuggestJobModel.portfolio_id == portfolio_id)
        related: list[dict[str, str]] = []
        for job in q.all():
            related.append(
                {
                    "entity_type": "ai_suggest_job",
                    "id": str(job.id),
                    "label": f"{job.trigger}:{job.status}",
                }
            )
            self.db.delete(job)
        return related

    def _purge_account(self, account: PayflowAccountModel) -> list[dict[str, str]]:
        related: list[dict[str, str]] = []
        comms = (
            self.db.query(PayflowCommunicationModel)
            .filter(PayflowCommunicationModel.account_id == account.id)
            .all()
        )
        for comm in comms:
            related.append(
                {"entity_type": "communication", "id": str(comm.id), "label": comm.code}
            )
            self.db.delete(comm)
        reviews = (
            self.db.query(PayflowHumanReviewModel)
            .filter(PayflowHumanReviewModel.account_id == account.id)
            .all()
        )
        for review in reviews:
            related.extend(self._purge_review(review))
        related.append(
            {
                "entity_type": "account",
                "id": str(account.id),
                "label": f"{account.customer_name} ({account.account_reference})",
            }
        )
        self.db.delete(account)
        return related

    def delete_portfolio(
        self,
        client_id: int,
        portfolio_id: int,
        *,
        actor: UserModel,
        source: str | None = None,
    ) -> dict:
        self._get_or_404(client_id)
        row = (
            self.db.query(PayflowPortfolioModel)
            .filter(
                PayflowPortfolioModel.id == portfolio_id,
                PayflowPortfolioModel.client_id == client_id,
            )
            .first()
        )
        if not row:
            raise NotFoundError("Portfolio not found")
        snapshot = snapshot_model(row)
        related: list[dict[str, str]] = []
        related.extend(self._purge_ai_suggest_jobs(portfolio_id=portfolio_id))
        accounts = (
            self.db.query(PayflowAccountModel)
            .filter(PayflowAccountModel.portfolio_id == portfolio_id)
            .all()
        )
        for account in accounts:
            related.extend(self._purge_account(account))
        strategies = (
            self.db.query(PayflowStrategyModel)
            .filter(PayflowStrategyModel.portfolio_id == portfolio_id)
            .all()
        )
        for strategy in strategies:
            related.append(
                {"entity_type": "workflow", "id": str(strategy.id), "label": strategy.name}
            )
            self.db.delete(strategy)
        record_deletion(
            self.db,
            actor=actor,
            module="payflow",
            entity_type="portfolio",
            entity_id=row.id,
            entity_label=f"{row.name} ({row.code})",
            source=source,
            record_snapshot=snapshot,
            related_deleted=related,
        )
        self.db.delete(row)
        self.db.commit()
        return {"message": "Portfolio deleted"}

    def delete_client(
        self, client_id: int, *, actor: UserModel, source: str | None = None
    ) -> dict:
        client = self._get_or_404(client_id)
        snapshot = snapshot_model(client)
        related: list[dict[str, str]] = []

        notifications = (
            self.db.query(PayflowNotificationModel)
            .filter(PayflowNotificationModel.client_id == client_id)
            .all()
        )
        for note in notifications:
            related.append({"entity_type": "notification", "id": str(note.id), "label": note.title})
            self.db.delete(note)

        assignments = (
            self.db.query(PayflowUserClientAssignmentModel)
            .filter(PayflowUserClientAssignmentModel.client_id == client_id)
            .all()
        )
        for assignment in assignments:
            perms = (
                self.db.query(PayflowUserClientPermissionModel)
                .filter(PayflowUserClientPermissionModel.assignment_id == assignment.id)
                .all()
            )
            for perm in perms:
                self.db.delete(perm)
            related.append(
                {
                    "entity_type": "supervisor_assignment",
                    "id": str(assignment.id),
                    "label": str(assignment.membership_id),
                }
            )
            self.db.delete(assignment)

        # Must go before portfolios — FK payflow_ai_suggest_jobs_portfolio_id_fkey
        related.extend(self._purge_ai_suggest_jobs(client_id=client_id))

        comms = (
            self.db.query(PayflowCommunicationModel)
            .filter(PayflowCommunicationModel.client_id == client_id)
            .all()
        )
        for comm in comms:
            related.append(
                {"entity_type": "communication", "id": str(comm.id), "label": comm.code}
            )
            self.db.delete(comm)

        reviews = (
            self.db.query(PayflowHumanReviewModel)
            .filter(PayflowHumanReviewModel.client_id == client_id)
            .all()
        )
        for review in reviews:
            related.extend(self._purge_review(review))

        accounts = (
            self.db.query(PayflowAccountModel)
            .filter(PayflowAccountModel.client_id == client_id)
            .all()
        )
        for account in accounts:
            # Comms/reviews for this client already removed above; remove any
            # leftover account-scoped rows then the account itself.
            related.extend(self._purge_account(account))

        strategies = (
            self.db.query(PayflowStrategyModel)
            .filter(PayflowStrategyModel.client_id == client_id)
            .all()
        )
        for strategy in strategies:
            related.append(
                {"entity_type": "workflow", "id": str(strategy.id), "label": strategy.name}
            )
            self.db.delete(strategy)

        portfolios = (
            self.db.query(PayflowPortfolioModel)
            .filter(PayflowPortfolioModel.client_id == client_id)
            .all()
        )
        for portfolio in portfolios:
            related.append(
                {"entity_type": "portfolio", "id": str(portfolio.id), "label": portfolio.name}
            )
            self.db.delete(portfolio)

        mappings = (
            self.db.query(PayflowClientFieldMappingModel)
            .filter(PayflowClientFieldMappingModel.client_id == client_id)
            .all()
        )
        for mapping in mappings:
            related.append(
                {
                    "entity_type": "field_mapping",
                    "id": str(mapping.id),
                    "label": mapping.source_field,
                }
            )
            self.db.delete(mapping)

        rules = (
            self.db.query(PayflowRuleModel)
            .filter(PayflowRuleModel.client_id == client_id)
            .all()
        )
        for rule in rules:
            related.append({"entity_type": "rule", "id": str(rule.id), "label": rule.name})
            self.db.delete(rule)

        record_deletion(
            self.db,
            actor=actor,
            module="payflow",
            entity_type="client",
            entity_id=client.id,
            entity_label=f"{client.name} ({client.code})",
            source=source,
            record_snapshot=snapshot,
            related_deleted=related,
        )
        self.db.delete(client)
        self.db.commit()
        return {"message": "Client deleted"}

    # ------------------------------------------------------------------
    # Bulk upload (CRM clients hierarchy + simple template)
    # ------------------------------------------------------------------

    def bulk_template_bytes(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Clients"
        ws.append(BULK_UPLOAD_HEADERS)
        # Minimal aligned examples only (header shape) — not product mock CRM data.
        examples = [
            [
                "2026-0001-001",
                "EX-MASTER",
                "Example Master Client",
                "EXMASTER",
                "True",
                "",
                "Consumer",
                "ops@example.com",
                "15155550100",
                "Jane Contact",
                "jane@example.com",
                "15155550100",
                "100 Main St",
                "Montreal",
                "Quebec",
                "Canada",
                "H3B1A1",
                "Financial Institutions",
                "Third Party",
                "English",
                "Canadian Dollars",
                "Enable",
            ],
            [
                "2026-0001-002",
                "EX-SUB-01",
                "Example Master Client",
                "EX Sub Book",
                "False",
                "2026-0001-001",
                "Loans",
                "ops@example.com",
                "15155550100",
                "",
                "",
                "",
                "100 Main St",
                "Montreal",
                "Quebec",
                "Canada",
                "H3B1A1",
                "Financial Institutions",
                "Third Party",
                "English",
                "Canadian Dollars",
                "Enable",
            ],
        ]
        phone_cols = {
            i + 1
            for i, h in enumerate(BULK_UPLOAD_HEADERS)
            if "phone" in h or h == "zip_code"
        }
        for row in examples:
            ws.append(row)
            r_idx = ws.max_row
            for c_idx in phone_cols:
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def bulk_validate(self, file_bytes: bytes, filename: str) -> dict:
        safe_name = (filename or "upload.xlsx").split("/")[-1]
        try:
            headers, data_rows = self._load_client_rows(file_bytes, filename)
        except ValidationAppError as exc:
            return self._bulk_preview_fail(safe_name, "File structure", exc.message)

        return self._process_client_hierarchy(
            headers, data_rows, filename=safe_name, dry_run=True, user=None
        )

    def bulk_upload(self, file_bytes: bytes, *, user: UserModel | None = None, filename: str = "upload.xlsx") -> dict:
        safe_name = (filename or "upload.xlsx").split("/")[-1]
        try:
            headers, data_rows = self._load_client_rows(file_bytes, filename)
        except ValidationAppError as exc:
            result = self._bulk_preview_fail(safe_name, "File structure", exc.message)
        else:
            result = self._process_client_hierarchy(
                headers, data_rows, filename=safe_name, dry_run=False, user=user, file_bytes=file_bytes
            )
        # Back-compat shape for older FE callers
        result["created_count"] = result.get("summary", {}).get("created", 0)
        result["error_count"] = max(
            result.get("summary", {}).get("failed", 0),
            result.get("summary", {}).get("rejected", 0),
        )
        result["created"] = [
            {
                "row": p.get("record_id"),
                "id": p.get("entity_id"),
                "code": p.get("sub_client") if p.get("is_sub") else p.get("client_code"),
                "name": p.get("sub_client_name") if p.get("is_sub") else p.get("client_name"),
            }
            for p in result.get("preview", [])
            if p.get("action") in ("Create", "Update") and p.get("entity_id")
        ]
        result["errors"] = [
            {
                "row": e.get("record_id"),
                "code": e.get("sub_client") if e.get("sub_client") not in (None, "—") else e.get("client"),
                "message": e.get("error"),
            }
            for e in result.get("errors", [])
        ]
        return result

    def _load_client_rows(
        self, file_bytes: bytes, filename: str
    ) -> tuple[list[str], list[dict[str, str]]]:
        lower = (filename or "").lower()
        if lower.endswith(".csv"):
            try:
                text = file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1")
            # excel dialect tolerates quoted commas so columns do not shift mid-row
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(
                io.StringIO(text), dialect=dialect, restkey="__overflow__"
            )
            if not reader.fieldnames:
                raise ValidationAppError("CSV file is empty or missing headers")
            headers = [str(h or "").strip().lower() for h in reader.fieldnames]
            if any(not h for h in headers):
                raise ValidationAppError(
                    "Header row has empty column name(s). Remove blank columns so "
                    "email / phone / contact stay aligned."
                )
            expected = len(headers)
            rows: list[dict[str, str]] = []
            for line_no, raw in enumerate(reader, start=2):
                overflow = raw.get("__overflow__")
                if overflow:
                    raise ValidationAppError(
                        f"Row {line_no} has more columns than the header "
                        f"({expected}). Check for an extra comma or a value "
                        "that shifted email/phone into the wrong column."
                    )
                row = {
                    str(k or "").strip().lower(): _cell_to_str(v)
                    for k, v in raw.items()
                    if k is not None and k != "__overflow__"
                }
                if any(row.values()):
                    rows.append(row)
            return headers, rows

        try:
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise ValidationAppError("Invalid Excel or CSV file") from exc
        ws = wb.active
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            raise ValidationAppError("Excel file is empty")
        headers = [_cell_to_str(h).lower() for h in raw_rows[0]]
        if any(not h for h in headers):
            # Allow trailing blank headers, but not blanks in the middle
            trimmed = list(headers)
            while trimmed and not trimmed[-1]:
                trimmed.pop()
            if any(not h for h in trimmed):
                raise ValidationAppError(
                    "Header row has empty column name(s). Remove blank columns so "
                    "email / phone / contact stay aligned."
                )
            headers = trimmed
        rows = []
        for row_num, raw in enumerate(raw_rows[1:], start=2):
            if not raw or all(c is None or _cell_to_str(c) == "" for c in raw):
                continue
            cells = list(raw)
            # Ignore trailing empty cells beyond header width
            while len(cells) > len(headers) and _cell_to_str(cells[-1]) == "":
                cells.pop()
            if len(cells) > len(headers):
                raise ValidationAppError(
                    f"Row {row_num} has more columns than the header "
                    f"({len(headers)}). A value is shifted — usually email/phone "
                    "or product landed in the wrong column."
                )
            row = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                val = cells[i] if i < len(cells) else None
                row[h] = _cell_to_str(val)
            if any(row.values()):
                rows.append(row)
        return headers, rows

    def _is_hierarchy_format(self, headers: list[str]) -> bool:
        return "is_master_client" in headers or (
            "client_number" in headers and "client_code" in headers
        )

    def _process_client_hierarchy(
        self,
        headers: list[str],
        data_rows: list[dict[str, str]],
        *,
        filename: str,
        dry_run: bool,
        user: UserModel | None,
        file_bytes: bytes | None = None,
    ) -> dict:
        if not data_rows:
            return self._bulk_preview_fail(
                filename,
                "File structure",
                "File has no data rows. Add at least one client record below the header row.",
            )

        if self._is_hierarchy_format(headers):
            missing = [h for h in CLIENT_IMPORT_MASTER_REQUIRED if h not in headers]
            if not any(h in headers for h in CLIENT_IMPORT_NAME_COLUMNS):
                missing.append("company_name or short_name")
            for _label, aliases in CLIENT_IMPORT_MASTER_PROFILE_HEADERS.items():
                if not any(a in headers for a in aliases):
                    missing.append(aliases[0])
            # De-dupe while preserving order
            seen: set[str] = set()
            missing_unique = []
            for m in missing:
                if m not in seen:
                    seen.add(m)
                    missing_unique.append(m)
            if missing_unique:
                return self._bulk_preview_fail(
                    filename,
                    "File structure",
                    f"Missing required column(s): {', '.join(missing_unique)}",
                )
            return self._process_crm_hierarchy(
                data_rows, filename=filename, dry_run=dry_run, user=user, file_bytes=file_bytes
            )

        # Legacy simple template: client_name + client_code
        if "client_name" not in headers or "client_code" not in headers:
            return self._bulk_preview_fail(
                filename,
                "File structure",
                "Expected CRM hierarchy columns (is_master_client, client_number, "
                "client_code) or legacy client_name/client_code columns.",
            )
        return self._process_simple_clients(
            data_rows, filename=filename, dry_run=dry_run, user=user, file_bytes=file_bytes
        )

    def _validate_master_profile_row(self, row: dict[str, str]) -> None:
        """Enforce Add Client mandatory fields on master rows."""
        client_type = _row_first(row, "client_type__name")
        if not client_type:
            raise ValidationAppError("Client type is required")

        contact_name = _row_first(row, "clientdemographiccontactinformation__full_name")
        if not contact_name:
            raise ValidationAppError("Primary contact name is required")

        email = _row_first(
            row,
            "email_address",
            "clientdemographiccontactinformation__email_address",
        )
        if not email:
            raise ValidationAppError("Primary contact email is required")
        if not _EMAIL_RE.match(email):
            raise ValidationAppError("Enter a valid primary contact email")

        phone = _row_first(
            row,
            "phone_number",
            "cell_number",
            "clientdemographiccontactinformation__phone_number",
        )
        if not phone:
            raise ValidationAppError("Primary contact phone is required")

        country = _row_first(row, "country__name")
        if not country:
            raise ValidationAppError("Country is required")

        language = _row_first(row, "correspondence_language__name")
        if not language:
            raise ValidationAppError("Language is required")

        _normalize_import_currency(_row_first(row, "currency__name"))

        postal = _row_first(row, "zip_code", "postal_code")
        if postal and not _is_valid_postal_code(postal):
            raise ValidationAppError(
                "Postal code must match format A1A 1A1 "
                "(letter-number-letter space number-letter-number)"
            )

        for date_field in _CLIENT_IMPORT_DATE_FIELDS:
            if date_field in row:
                _validate_optional_import_date(row.get(date_field) or "", date_field)

        # Optional boolean columns — reject garbage when present.
        # Note: CRM "sliding_scale" is NOT a boolean (tier JSON/list); do not validate it here.
        for bool_field in ("restricted_access_client",):
            raw = row.get(bool_field)
            if raw is None or str(raw).strip() == "":
                continue
            _parse_import_bool(raw, field=bool_field)

    def _process_crm_hierarchy(
        self,
        data_rows: list[dict[str, str]],
        *,
        filename: str,
        dry_run: bool,
        user: UserModel | None,
        file_bytes: bytes | None,
    ) -> dict:
        preview_masters: list[dict] = []
        preview_subs: list[dict] = []
        errors: list[dict] = []
        created = updated = unchanged = failed = 0
        new_clients = existing_clients = new_sub_clients = existing_sub_clients = 0
        masters_in_file: dict[str, dict[str, str]] = {}
        seen_master_codes: set[str] = set()
        seen_sub_codes: set[tuple[str, str]] = set()
        seen_numbers: set[str] = set()
        flag_error_rows: set[int] = set()

        # Pass 1 — masters
        for row_num, row in enumerate(data_rows, start=2):
            record_id = row.get("client_number") or f"Row {row_num}"
            code = _normalize_code(row.get("client_code") or "")
            number = (row.get("client_number") or "").strip()
            name = (row.get("company_name") or row.get("short_name") or "").strip()
            try:
                is_master = _parse_import_bool(
                    row.get("is_master_client"), field="is_master_client"
                )
            except ValidationAppError as exc:
                failed += 1
                flag_error_rows.add(row_num)
                field = _guess_client_import_field(exc.message)
                errors.append(
                    {
                        "record_id": record_id or f"Row {row_num}",
                        "client": name or code or "—",
                        "sub_client": "—",
                        "field": field,
                        "error": exc.message,
                        "status": "Rejected",
                    }
                )
                preview_masters.append(
                    {
                        "id": f"master-err-{row_num}",
                        "record_id": record_id or f"Row {row_num}",
                        "client": code or "—",
                        "client_name": name or code or f"Row {row_num}",
                        "client_code": code or "—",
                        "sub_client": "—",
                        "sub_client_name": "—",
                        "action": "Error",
                        "note": exc.message,
                        "is_sub": False,
                        "entity_id": None,
                    }
                )
                continue
            if not is_master:
                continue
            try:
                parent_ref = (row.get("master_client__client_number") or "").strip()
                if parent_ref:
                    raise ValidationAppError(
                        "Master clients cannot have master_client__client_number set. "
                        "Clear that column for master rows, or set is_master_client to "
                        "False for sub-clients."
                    )
                if not number:
                    raise ValidationAppError("client_number is required for master clients")
                if not code:
                    raise ValidationAppError("client_code is required")
                if not name:
                    raise ValidationAppError("company_name or short_name is required")
                self._validate_master_profile_row(row)
                if number.lower() in seen_numbers:
                    raise ConflictError(
                        f'Duplicate client_number "{number}" in this file. '
                        "Each client_number must be unique."
                    )
                if code.lower() in seen_master_codes:
                    raise ConflictError(
                        f'Duplicate client_code "{code}" in this file. '
                        "Each master client_code must be unique."
                    )
                self._assert_master_number_not_used_as_sub(number)

                payload = self._crm_row_to_client_payload(row, data_source="crm")
                existing = self._resolve_client_for_upsert(crm_number=number, code=code)
                entity_id = existing.id if existing else None
                if existing and self._client_crm_unchanged(existing, payload):
                    action = "No Change"
                    unchanged += 1
                    existing_clients += 1
                elif existing:
                    action = "Update"
                    if not dry_run:
                        client, _act = self.upsert_client(
                            {
                                **payload,
                                "added_through": PayflowClientAddedThrough.FILE_UPLOAD.value,
                            },
                            commit=False,
                            soft_client_type=True,
                        )
                        client.connection_status = PayflowConnectionStatus.CONNECTED.value
                        entity_id = client.id
                    updated += 1
                    existing_clients += 1
                else:
                    action = "Create"
                    if not dry_run:
                        client, _act = self.upsert_client(
                            {
                                **payload,
                                "added_through": PayflowClientAddedThrough.FILE_UPLOAD.value,
                            },
                            commit=False,
                            soft_client_type=True,
                        )
                        client.connection_status = PayflowConnectionStatus.CONNECTED.value
                        entity_id = client.id
                    created += 1
                    new_clients += 1

                has_subs_in_file = any(
                    not _truthy(r.get("is_master_client"))
                    and (r.get("master_client__client_number") or "").strip().lower()
                    == number.lower()
                    for r in data_rows
                )
                setup_issues: list[str] = []
                if entity_id and not dry_run:
                    live = (
                        self.db.query(PayflowClientModel)
                        .filter(PayflowClientModel.id == entity_id)
                        .first()
                    )
                    if live is not None:
                        setup_issues = self._setup_issue_messages(live)
                        if has_subs_in_file:
                            setup_issues = [
                                i
                                for i in setup_issues
                                if "sub-clients / portfolios" not in i.lower()
                            ]
                elif existing is not None:
                    setup_issues = self._setup_issue_messages(existing)
                    if has_subs_in_file:
                        setup_issues = [
                            i
                            for i in setup_issues
                            if "sub-clients / portfolios" not in i.lower()
                        ]
                else:
                    setup_issues = self._simulate_setup_issues_after_file_import(
                        payload, has_subs_in_file=has_subs_in_file
                    )

                # Only accept identity after validation/create path succeeds
                seen_numbers.add(number.lower())
                seen_master_codes.add(code.lower())
                masters_in_file[number.lower()] = row
                note = number
                if setup_issues:
                    note = f"{number} · Setup incomplete: {'; '.join(setup_issues)}"
                    errors.append(
                        {
                            "record_id": record_id,
                            "client": name or code or "—",
                            "sub_client": "—",
                            "field": "setup",
                            "error": (
                                "Client will import as Draft with remaining setup: "
                                + "; ".join(setup_issues)
                            ),
                            "status": "Attention",
                        }
                    )
                preview_masters.append(
                    {
                        "id": f"master-{row_num}",
                        "record_id": record_id,
                        "client": code,
                        "client_name": name,
                        "client_code": code,
                        "sub_client": "—",
                        "sub_client_name": "—",
                        "action": action,
                        "note": note,
                        "is_sub": False,
                        "entity_id": entity_id,
                        "setup_issues": setup_issues,
                    }
                )
            except (ValidationAppError, ConflictError) as exc:
                failed += 1
                field = _guess_client_import_field(exc.message)
                errors.append(
                    {
                        "record_id": record_id or f"Row {row_num}",
                        "client": name or code or "—",
                        "sub_client": "—",
                        "field": field,
                        "error": exc.message,
                        "status": "Skipped" if "Duplicate" in exc.message else "Rejected",
                    }
                )
                preview_masters.append(
                    {
                        "id": f"master-err-{row_num}",
                        "record_id": record_id or f"Row {row_num}",
                        "client": code or "—",
                        "client_name": name or code or f"Row {row_num}",
                        "client_code": code or "—",
                        "sub_client": "—",
                        "sub_client_name": "—",
                        "action": "Error",
                        "note": exc.message,
                        "is_sub": False,
                        "entity_id": None,
                    }
                )

        # Resolve masters already in DB for children whose parent is not in this file
        db_masters_by_number = {
            (c.crm_client_number or "").strip().lower(): c
            for c in self.db.query(PayflowClientModel).all()
            if (c.crm_client_number or "").strip()
        }

        # Pass 2 — sub-clients
        for row_num, row in enumerate(data_rows, start=2):
            if row_num in flag_error_rows:
                continue
            record_id = row.get("client_number") or f"Row {row_num}"
            code = _normalize_code(row.get("client_code") or "")
            number = (row.get("client_number") or "").strip()
            master_number = (row.get("master_client__client_number") or "").strip()
            name = (
                row.get("short_name")
                or row.get("product")
                or row.get("company_name")
                or ""
            ).strip()
            try:
                if _parse_import_bool(
                    row.get("is_master_client"), field="is_master_client"
                ):
                    continue
                if not master_number:
                    raise ValidationAppError(
                        "Master client number is required for sub-clients. "
                        "Set master_client__client_number to the parent client_number."
                    )
                if not code:
                    raise ValidationAppError("client_code is required")
                if not name:
                    raise ValidationAppError("short_name, product, or company_name is required")
                if number and number.lower() in seen_numbers:
                    raise ConflictError(
                        f'Duplicate client_number "{number}" in this file. '
                        "Each client_number must be unique."
                    )
                sub_code_key = (master_number.lower(), code.lower())
                if sub_code_key in seen_sub_codes:
                    raise ConflictError(
                        f'Duplicate client_code "{code}" under master '
                        f'"{master_number}" in this file. Sub-client codes must be '
                        "unique within the same master."
                    )

                parent = self._find_client(crm_number=master_number, code=None)
                if parent is None:
                    parent = db_masters_by_number.get(master_number.lower())
                master_row = masters_in_file.get(master_number.lower())
                if parent is None and master_row is None:
                    raise ValidationAppError(
                        f'Master client "{master_number}" was not found. '
                        "Add the master row first, or fix master_client__client_number."
                    )

                self._assert_sub_import_identity(
                    number=number or None,
                    master_number=master_number,
                    parent=parent,
                )

                parent_name = (
                    parent.name
                    if parent
                    else (
                        master_row.get("company_name")
                        or master_row.get("short_name")
                        or master_number
                    )
                )
                parent_code = (
                    parent.code
                    if parent
                    else _normalize_code(master_row.get("client_code") or "")
                )

                self._assert_sub_belongs_to_master(
                    number=number or None,
                    code=code,
                    master_number=master_number,
                    master_code=parent_code,
                )

                existing = (
                    self._resolve_portfolio_for_upsert(
                        parent.id, crm_number=number or None, code=code
                    )
                    if parent
                    else None
                )
                description = (row.get("product") or "").strip() or None
                entity_id = existing.id if existing else None
                if existing and self._portfolio_crm_unchanged(
                    existing, name=name, code=code, crm_number=number or None, description=description
                ):
                    action = "No Change"
                    unchanged += 1
                    existing_sub_clients += 1
                elif existing:
                    action = "Update"
                    if not dry_run:
                        if parent is None:
                            parent = self._find_client(crm_number=master_number, code=None)
                        if parent is None:
                            raise ValidationAppError(
                                f'Master client "{master_number}" was not found. '
                                "Add the master row first, or fix master_client__client_number."
                            )
                        port, _act = self.upsert_portfolio(
                            parent.id,
                            {
                                "name": name,
                                "code": code,
                                "crm_client_number": number or None,
                                "description": description,
                                "status": PayflowPortfolioStatus.ONBOARDING.value,
                            },
                            commit=False,
                        )
                        entity_id = port.id
                        db_masters_by_number[master_number.lower()] = parent
                        parent_name = parent.name
                        parent_code = parent.code
                    updated += 1
                    existing_sub_clients += 1
                else:
                    action = "Create"
                    if not dry_run:
                        if parent is None:
                            parent = self._find_client(crm_number=master_number, code=None)
                        if parent is None:
                            raise ValidationAppError(
                                f'Master client "{master_number}" was not found. '
                                "Add the master row first, or fix master_client__client_number."
                            )
                        port, _act = self.upsert_portfolio(
                            parent.id,
                            {
                                "name": name,
                                "code": code,
                                "crm_client_number": number or None,
                                "description": description,
                                "status": PayflowPortfolioStatus.ONBOARDING.value,
                            },
                            commit=False,
                        )
                        entity_id = port.id
                        db_masters_by_number[master_number.lower()] = parent
                        parent_name = parent.name
                        parent_code = parent.code
                    created += 1
                    new_sub_clients += 1
                if number:
                    seen_numbers.add(number.lower())
                seen_sub_codes.add((master_number.lower(), code.lower()))
                preview_subs.append(
                    {
                        "id": f"sub-{row_num}",
                        "record_id": record_id,
                        "client": parent_code or "—",
                        "client_name": parent_name,
                        "client_code": parent_code or "—",
                        "sub_client": code,
                        "sub_client_name": name,
                        "action": action,
                        "note": number or code,
                        "is_sub": True,
                        "entity_id": entity_id,
                    }
                )
            except (ValidationAppError, ConflictError) as exc:
                failed += 1
                field = _guess_client_import_field(exc.message)
                errors.append(
                    {
                        "record_id": record_id or f"Row {row_num}",
                        "client": master_number or "—",
                        "sub_client": code or name or "—",
                        "field": field,
                        "error": exc.message,
                        "status": "Skipped" if "Duplicate" in exc.message else "Rejected",
                    }
                )
                preview_subs.append(
                    {
                        "id": f"sub-err-{row_num}",
                        "record_id": record_id or f"Row {row_num}",
                        "client": master_number or "—",
                        "client_name": master_number or "—",
                        "client_code": master_number or "—",
                        "sub_client": code or "—",
                        "sub_client_name": name or code or f"Row {row_num}",
                        "action": "Error",
                        "note": exc.message,
                        "is_sub": True,
                        "entity_id": None,
                    }
                )

        preview = self._assemble_client_hierarchy_preview(preview_masters, preview_subs)
        rejected = sum(1 for e in errors if (e.get("status") or "").lower() == "rejected")
        successful = created + updated + unchanged
        summary = {
            "total": created + updated + unchanged + failed,
            "created": created,
            "updated": updated,
            "unchanged": unchanged,
            "failed": failed,
            "rejected": rejected,
            "successful": successful,
            "new_clients": new_clients,
            "existing_clients": existing_clients,
            "new_sub_clients": new_sub_clients,
            "existing_sub_clients": existing_sub_clients,
        }
        setup_message = None
        if new_clients > 0:
            setup_message = (
                f"{new_clients} new client(s) import as Draft. "
                "Open each client and use Complete Setup to finish onboarding "
                "(supervisors, branding, and remaining configuration)."
            )

        if not dry_run:
            run = self._persist_client_import_run(
                user=user,
                filename=filename,
                file_bytes=file_bytes,
                created=created,
                updated=updated,
                unchanged=unchanged,
                failed=failed,
                errors=errors,
            )
            self.db.commit()
            return {
                "ok": True,
                "file_name": filename,
                "status": run.status if run else ("Completed with Errors" if failed else "Completed"),
                "import_id": run.id if run else None,
                "message": setup_message,
                "summary": summary,
                "preview": preview,
                "errors": errors,
            }

        # Validate/preview: fail the file when nothing is importable.
        if successful == 0:
            return {
                "ok": False,
                "file_name": filename,
                "status": "Failed",
                "message": (
                    errors[0]["error"]
                    if len(errors) == 1
                    else "No valid client records found. Fix the errors below and upload again."
                ),
                "summary": summary,
                "preview": preview,
                "errors": errors
                or [
                    {
                        "record_id": "File",
                        "client": "—",
                        "sub_client": "—",
                        "field": "File structure",
                        "error": "No valid client records found in the file.",
                        "status": "Rejected",
                    }
                ],
            }

        return {
            "ok": True,
            "file_name": filename,
            "status": "Ready",
            "message": setup_message,
            "summary": summary,
            "preview": preview,
            "errors": errors,
        }

    @staticmethod
    def _assemble_client_hierarchy_preview(
        masters: list[dict],
        subs: list[dict],
        *,
        limit: int = 120,
    ) -> list[dict]:
        """Build a hierarchical preview: parent then its sub-clients (not masters-only)."""
        subs_by_parent: dict[str, list[dict]] = {}
        for s in subs:
            key = (s.get("client_name") or "").strip().lower()
            if not key:
                key = (s.get("client_code") or s.get("client") or "").strip().lower()
            subs_by_parent.setdefault(key, []).append(s)

        with_subs = [
            m
            for m in masters
            if (m.get("client_name") or "").strip().lower() in subs_by_parent
        ]
        alone = [
            m
            for m in masters
            if (m.get("client_name") or "").strip().lower() not in subs_by_parent
        ]
        ordered_masters = with_subs + alone

        preview: list[dict] = []
        included_sub_ids: set[str] = set()
        for m in ordered_masters:
            if len(preview) >= limit:
                break
            preview.append(m)
            key = (m.get("client_name") or "").strip().lower()
            for s in subs_by_parent.get(key, []):
                if len(preview) >= limit:
                    break
                preview.append(s)
                included_sub_ids.add(str(s.get("id")))

        for s in subs:
            if len(preview) >= limit:
                break
            if str(s.get("id")) in included_sub_ids:
                continue
            preview.append(s)
        return preview

    def _process_simple_clients(
        self,
        data_rows: list[dict[str, str]],
        *,
        filename: str,
        dry_run: bool,
        user: UserModel | None,
        file_bytes: bytes | None,
    ) -> dict:
        preview: list[dict] = []
        errors: list[dict] = []
        created = updated = unchanged = failed = 0
        seen_codes: set[str] = set()
        for row_num, row in enumerate(data_rows, start=2):
            name = (row.get("client_name") or row.get("company_name") or "").strip()
            code = _normalize_code(row.get("client_code") or "")
            record_id = f"Row {row_num}"
            try:
                if not name:
                    raise ValidationAppError("Client name is required")
                if not code:
                    raise ValidationAppError("Client code / reference is required")
                if code.lower() in seen_codes:
                    raise ConflictError(f"Duplicate client code in file: {code}")
                seen_codes.add(code.lower())
                existing = self._find_client(crm_number=None, code=code)
                payload = {
                    "name": name,
                    "code": code,
                    "client_type": row.get("client_type") or "Third Party",
                    "business_domain": row.get("business_domain") or "Collections",
                    "industry": row.get("industry") or None,
                    "ai_mode": row.get("ai_mode") or "Supervised AI",
                    "data_source_type": "file",
                    "added_through": PayflowClientAddedThrough.FILE_UPLOAD.value,
                }
                entity_id = existing.id if existing else None
                if existing and (existing.name or "").strip() == name and (existing.code or "").lower() == code.lower():
                    action = "No Change"
                    unchanged += 1
                elif existing:
                    action = "Update"
                    if not dry_run:
                        client, _act = self.upsert_client(
                            payload, commit=False, soft_client_type=True
                        )
                        entity_id = client.id
                    updated += 1
                else:
                    action = "Create"
                    if not dry_run:
                        client, _act = self.upsert_client(
                            payload, commit=False, soft_client_type=True
                        )
                        entity_id = client.id
                    created += 1
                if len(preview) < 50:
                    preview.append(
                        {
                            "id": f"row-{row_num}",
                            "record_id": record_id,
                            "client": name,
                            "client_name": name,
                            "client_code": code,
                            "sub_client": "—",
                            "sub_client_name": "—",
                            "action": action,
                            "note": code,
                            "is_sub": False,
                            "entity_id": entity_id,
                        }
                    )
            except (ValidationAppError, ConflictError) as exc:
                failed += 1
                field = _guess_client_import_field(exc.message)
                errors.append(
                    {
                        "record_id": record_id,
                        "client": name or "—",
                        "sub_client": "—",
                        "field": field,
                        "error": exc.message,
                        "status": "Skipped" if "Duplicate" in exc.message else "Rejected",
                    }
                )
                if len(preview) < 50:
                    preview.append(
                        {
                            "id": f"row-err-{row_num}",
                            "record_id": record_id,
                            "client": name or code or "—",
                            "client_name": name or code or record_id,
                            "client_code": code or "—",
                            "sub_client": "—",
                            "sub_client_name": "—",
                            "action": "Error",
                            "note": exc.message,
                            "is_sub": False,
                            "entity_id": None,
                        }
                    )

        rejected = sum(1 for e in errors if (e.get("status") or "").lower() == "rejected")
        successful = created + updated + unchanged
        summary = {
            "total": created + updated + unchanged + failed,
            "created": created,
            "updated": updated,
            "unchanged": unchanged,
            "failed": failed,
            "rejected": rejected,
            "successful": successful,
            "new_clients": created,
            "existing_clients": updated + unchanged,
            "new_sub_clients": 0,
            "existing_sub_clients": 0,
        }

        if not dry_run:
            run = self._persist_client_import_run(
                user=user,
                filename=filename,
                file_bytes=file_bytes,
                created=created,
                updated=updated,
                unchanged=unchanged,
                failed=failed,
                errors=errors,
            )
            self.db.commit()
            return {
                "ok": True,
                "file_name": filename,
                "status": run.status if run else ("Completed with Errors" if failed else "Completed"),
                "import_id": run.id if run else None,
                "message": None,
                "summary": summary,
                "preview": preview,
                "errors": errors,
            }

        if successful == 0:
            return {
                "ok": False,
                "file_name": filename,
                "status": "Failed",
                "message": (
                    errors[0]["error"]
                    if len(errors) == 1
                    else "No valid client records found. Fix the errors below and upload again."
                ),
                "summary": summary,
                "preview": preview,
                "errors": errors
                or [
                    {
                        "record_id": "File",
                        "client": "—",
                        "sub_client": "—",
                        "field": "File structure",
                        "error": "No valid client records found in the file.",
                        "status": "Rejected",
                    }
                ],
            }

        return {
            "ok": True,
            "file_name": filename,
            "status": "Ready",
            "message": None,
            "summary": summary,
            "preview": preview,
            "errors": errors,
        }

    def _crm_row_to_client_payload(self, row: dict[str, str], *, data_source: str) -> dict[str, Any]:
        email = (
            row.get("email_address")
            or row.get("clientdemographiccontactinformation__email_address")
            or ""
        ).strip()
        phone = (
            row.get("phone_number")
            or row.get("cell_number")
            or row.get("clientdemographiccontactinformation__phone_number")
            or ""
        ).strip()
        contact_name = (
            row.get("clientdemographiccontactinformation__full_name") or ""
        ).strip()
        contact_title = (
            row.get("clientdemographiccontactinformation__contact_title") or ""
        ).strip()
        number = (row.get("client_number") or "").strip()
        return {
            "name": (row.get("company_name") or row.get("short_name") or "").strip(),
            "code": _normalize_code(row.get("client_code") or ""),
            "short_name": (row.get("short_name") or "").strip() or None,
            "brand_name": (row.get("short_name") or "").strip() or None,
            "industry": (row.get("client_industry__name") or "").strip() or None,
            "client_type": row.get("client_type__name") or "Third Party",
            "business_domain": "Collections",
            "ai_mode": "Supervised AI",
            "data_source_type": data_source,
            "crm_client_number": number or None,
            "integration_ref": number or None,
            "crm_system_name": "CRM",
            "contact_name": contact_name or None,
            "contact_title": contact_title or None,
            "contact_email": email or None,
            "contact_phone": phone or None,
            "address_line1": (row.get("address_line_1") or "").strip() or None,
            "address_line2": (row.get("address_line_2") or "").strip() or None,
            "city": (row.get("city__name") or "").strip() or None,
            "province_state": (row.get("province__name") or "").strip() or None,
            "country": (row.get("country__name") or "").strip() or None,
            "postal_code": _normalize_postal_code(row.get("zip_code") or "") or None,
            "correspondence_language": (
                row.get("correspondence_language__name") or ""
            ).strip()
            or None,
            "currency_code": _normalize_import_currency(
                (row.get("currency__name") or "").strip()
            )
            if (row.get("currency__name") or "").strip()
            else None,
            "crm_status": (row.get("client_status__name") or "").strip() or None,
        }

    def _persist_client_import_run(
        self,
        *,
        user: UserModel | None,
        filename: str,
        file_bytes: bytes | None,
        created: int,
        updated: int,
        unchanged: int = 0,
        failed: int,
        errors: list[dict],
    ) -> PayflowImportRunModel | None:
        stored_path = None
        if file_bytes:
            _CLIENT_UPLOADS.mkdir(parents=True, exist_ok=True)
            safe = Path(filename).name
            dest = _CLIENT_UPLOADS / f"{int(_utcnow().timestamp())}-{safe}"
            dest.write_bytes(file_bytes)
            try:
                stored_path = str(dest.relative_to(_BE_ROOT)).replace("\\", "/")
            except ValueError:
                stored_path = str(dest)

        # Status lifecycle: Uploaded → Validating → Processing → terminal
        # (client hierarchy is processed synchronously before persist; record Uploaded then terminal.)
        if created or updated or unchanged:
            status = "Completed with Errors" if failed else "Completed"
        else:
            status = "Failed"

        run = PayflowImportRunModel(
            kind="client",
            file_name=filename,
            stored_path=stored_path,
            uploaded_by_user_id=user.id if user else None,
            uploaded_by_name=(user.full_name if user else None) or (user.email if user else "System"),
            status="Uploaded",
            total_count=created + updated + unchanged + failed,
            created_count=created,
            updated_count=updated,
            unchanged_count=unchanged,
            failed_count=failed,
        )
        self.db.add(run)
        self.db.flush()
        run.status = "Validating"
        self.db.flush()
        run.status = "Processing"
        self.db.flush()
        run.status = status
        run.completed_at = _utcnow()
        for err in errors:
            self.db.add(
                PayflowImportErrorModel(
                    run_id=run.id,
                    record_id=str(err.get("record_id") or "—"),
                    client=str(err.get("client") or "—"),
                    sub_client=str(err.get("sub_client") or "—"),
                    field=str(err.get("field") or "—"),
                    message=str(err.get("error") or "Error"),
                    status=str(err.get("status") or "Rejected"),
                )
            )
        return run

    @staticmethod
    def _client_crm_unchanged(client: PayflowClientModel, payload: dict[str, Any]) -> bool:
        name = (payload.get("name") or "").strip()
        code = _normalize_code(payload.get("code") or "")
        crm_number = (payload.get("crm_client_number") or "").strip() or None
        industry = (payload.get("industry") or payload.get("category") or "").strip() or None
        brand = (payload.get("brand_name") or payload.get("short_name") or "").strip() or ""
        if (client.name or "").strip() != name:
            return False
        if (client.code or "").lower() != code.lower():
            return False
        if (client.crm_client_number or None) != crm_number:
            return False
        if (client.category or None) != industry:
            return False
        if brand and (client.brand_name or "") != brand:
            return False
        for field in _PRIMARY_CONTACT_FIELDS:
            if field not in payload:
                continue
            incoming = (str(payload[field]).strip() if payload.get(field) is not None else None) or None
            current = getattr(client, field, None) or None
            if isinstance(current, str):
                current = current.strip() or None
            if incoming != current:
                return False
        return True

    @staticmethod
    def _portfolio_crm_unchanged(
        row: PayflowPortfolioModel,
        *,
        name: str,
        code: str,
        crm_number: str | None,
        description: str | None,
    ) -> bool:
        if (row.name or "").strip() != name:
            return False
        if (row.code or "").lower() != code.lower():
            return False
        if (row.crm_client_number or None) != (crm_number or None):
            return False
        desc = (description or "").strip() or None
        if (row.description or None) != desc:
            return False
        return True

    @staticmethod
    def _bulk_preview_fail(file_name: str, field: str, message: str) -> dict:
        return {
            "ok": False,
            "file_name": file_name,
            "status": "Failed",
            "message": message,
            "summary": {
                "total": 0,
                "created": 0,
                "updated": 0,
                "unchanged": 0,
                "failed": 0,
                "rejected": 1,
                "successful": 0,
                "new_clients": 0,
                "existing_clients": 0,
                "new_sub_clients": 0,
                "existing_sub_clients": 0,
            },
            "preview": [],
            "errors": [
                {
                    "record_id": "File",
                    "client": "—",
                    "sub_client": "—",
                    "field": field,
                    "error": message,
                    "status": "Rejected",
                }
            ],
        }

    # ------------------------------------------------------------------
    # Internals
    # ------------------------------------------------------------------

    def _get_or_404(self, client_id: int) -> PayflowClientModel:
        client = (
            self.db.query(PayflowClientModel)
            .options(
                joinedload(PayflowClientModel.field_mappings),
                joinedload(PayflowClientModel.portfolios).joinedload(PayflowPortfolioModel.accounts),
                joinedload(PayflowClientModel.portfolios).joinedload(PayflowPortfolioModel.strategies),
                joinedload(PayflowClientModel.assignments)
                .joinedload(PayflowUserClientAssignmentModel.membership)
                .joinedload(PayflowUserMembershipModel.user),
                joinedload(PayflowClientModel.assignments)
                .joinedload(PayflowUserClientAssignmentModel.membership)
                .joinedload(PayflowUserMembershipModel.role),
                joinedload(PayflowClientModel.assignments)
                .joinedload(PayflowUserClientAssignmentModel.permissions)
                .joinedload(PayflowUserClientPermissionModel.permission),
            )
            .filter(PayflowClientModel.id == client_id)
            .first()
        )
        if not client:
            raise NotFoundError("Client not found")
        return client

    def _assert_unique_code(self, code: str, exclude_id: int | None = None) -> None:
        q = self.db.query(PayflowClientModel).filter(PayflowClientModel.code == code)
        if exclude_id is not None:
            q = q.filter(PayflowClientModel.id != exclude_id)
        if q.first():
            raise ConflictError("Client code / reference already exists")

    def _assert_unique_crm_number(self, crm_number: str, exclude_id: int | None = None) -> None:
        q = self.db.query(PayflowClientModel).filter(
            PayflowClientModel.crm_client_number == crm_number
        )
        if exclude_id is not None:
            q = q.filter(PayflowClientModel.id != exclude_id)
        if q.first():
            raise ConflictError("CRM client number already linked to another client")

    @staticmethod
    def _assert_crm_number_immutable(
        *,
        existing: str | None,
        incoming: str | None,
        code: str | None = None,
        entity: str = "client",
    ) -> None:
        """Reject attempts to rewrite a locked client_number / sub-client number."""
        current = (existing or "").strip()
        new = (incoming or "").strip()
        if not current:
            return
        who = "client" if entity == "client" else "sub-client"
        if not new:
            raise ValidationAppError(
                f"The {who} number \"{current}\" cannot be removed"
                + (f' for client_code "{code}"' if code else "")
                + ". Keep the original client_number."
            )
        if new.lower() != current.lower():
            raise ValidationAppError(
                f'The {who} number cannot be changed from "{current}" to "{new}"'
                + (f' for client_code "{code}"' if code else "")
                + ". Keep the original number, or use a new number and a new "
                "code to create a new record."
            )

    @staticmethod
    def _assert_sub_belongs_to_master(
        *,
        number: str | None,
        code: str,
        master_number: str,
        master_code: str,
    ) -> None:
        """Reject subs whose number/code do not follow the parent master pattern."""
        master_number = (master_number or "").strip()
        master_code = _normalize_code(master_code or "")
        number = (number or "").strip()
        code = _normalize_code(code or "")
        if not master_number or not master_code:
            return

        num_ok = bool(number) and number.lower().startswith(
            f"{master_number.lower()}-"
        )
        code_ok = bool(code) and code.lower().startswith(f"{master_code.lower()}-")
        if num_ok and code_ok:
            return
        raise ValidationAppError(
            f'This sub-client does not belong to master "{master_number}" / '
            f'"{master_code}". Use a client_number starting with '
            f'"{master_number}-" and a client_code starting with "{master_code}-".'
        )

    def _find_portfolio_by_crm_number_global(
        self, crm_number: str
    ) -> PayflowPortfolioModel | None:
        number = (crm_number or "").strip()
        if not number:
            return None
        return (
            self.db.query(PayflowPortfolioModel)
            .filter(PayflowPortfolioModel.crm_client_number == number)
            .first()
        )

    def _assert_master_number_not_used_as_sub(self, crm_number: str) -> None:
        """Reject master rows whose client_number already belongs to a sub-client."""
        number = (crm_number or "").strip()
        if not number:
            return
        port = self._find_portfolio_by_crm_number_global(number)
        if port is None:
            return
        raise ValidationAppError(
            f'Client number "{number}" is already used by a sub-client'
            + (f' (client_code "{port.code}")' if port.code else "")
            + ". A master client cannot reuse a sub-client number."
        )

    def _assert_sub_import_identity(
        self,
        *,
        number: str | None,
        master_number: str,
        parent: PayflowClientModel | None,
    ) -> None:
        """Lock sub parent relationship and block master/sub role flips."""
        num = (number or "").strip() or None
        if not num:
            return

        as_master = self._find_client(crm_number=num, code=None)
        if as_master is not None:
            raise ValidationAppError(
                f'Client number "{num}" is already used by a master client'
                + (f' (client_code "{as_master.code}")' if as_master.code else "")
                + ". A sub-client cannot reuse a master client number."
            )

        global_port = self._find_portfolio_by_crm_number_global(num)
        if global_port is None:
            return

        existing_parent = (
            self.db.query(PayflowClientModel)
            .filter(PayflowClientModel.id == global_port.client_id)
            .first()
        )
        existing_master = (
            (existing_parent.crm_client_number or "").strip()
            if existing_parent
            else ""
        ) or f"client #{global_port.client_id}"

        if parent is not None:
            if global_port.client_id != parent.id:
                raise ValidationAppError(
                    f'Sub-client "{num}" already belongs to master '
                    f'"{existing_master}". It cannot be moved under '
                    f'"{master_number}".'
                )
            return

        # Parent only present in this file (not flushed yet) — still block reparent
        # when the sub already exists under a different master in DB.
        if existing_master.lower() != (master_number or "").strip().lower():
            raise ValidationAppError(
                f'Sub-client "{num}" already belongs to master '
                f'"{existing_master}". It cannot be moved under '
                f'"{master_number}".'
            )

    def _resolve_client_for_upsert(
        self, *, crm_number: str | None, code: str | None
    ) -> PayflowClientModel | None:
        """Match by client_number first. Code fallback never rewrites identity.

        New client_number + new client_code → no match → caller creates.
        Same client_code + different client_number → error.
        """
        number = (crm_number or "").strip() or None
        code_norm = _normalize_code(code or "") if code else ""
        if number:
            by_number = (
                self.db.query(PayflowClientModel)
                .filter(PayflowClientModel.crm_client_number == number)
                .first()
            )
            if by_number:
                return by_number
        if code_norm:
            by_code = (
                self.db.query(PayflowClientModel)
                .filter(PayflowClientModel.code == code_norm)
                .first()
            )
            if by_code is None:
                return None
            existing_number = (by_code.crm_client_number or "").strip()
            if number and existing_number and existing_number.lower() != number.lower():
                raise ValidationAppError(
                    f'Client number cannot be changed for an existing client_code '
                    f'"{code_norm}". Keep the original number "{existing_number}", '
                    "or use a new number and a new code to create a new client."
                )
            return by_code
        return None

    def _find_client(
        self, *, crm_number: str | None, code: str | None
    ) -> PayflowClientModel | None:
        if crm_number:
            row = (
                self.db.query(PayflowClientModel)
                .filter(PayflowClientModel.crm_client_number == crm_number)
                .first()
            )
            if row:
                return row
        if code:
            return (
                self.db.query(PayflowClientModel)
                .filter(PayflowClientModel.code == code)
                .first()
            )
        return None

    def _resolve_portfolio_for_upsert(
        self,
        client_id: int,
        *,
        crm_number: str | None,
        code: str | None,
    ) -> PayflowPortfolioModel | None:
        """Match sub-client by client_number first; never rewrite via code."""
        number = (crm_number or "").strip() or None
        code_norm = _normalize_code(code or "") if code else ""
        if number:
            as_master = self._find_client(crm_number=number, code=None)
            if as_master is not None:
                raise ValidationAppError(
                    f'Client number "{number}" is already used by a master client'
                    + (f' (client_code "{as_master.code}")' if as_master.code else "")
                    + ". A sub-client cannot reuse a master client number."
                )
            global_row = self._find_portfolio_by_crm_number_global(number)
            if global_row is not None and global_row.client_id != client_id:
                existing_parent = (
                    self.db.query(PayflowClientModel)
                    .filter(PayflowClientModel.id == global_row.client_id)
                    .first()
                )
                existing_master = (
                    (existing_parent.crm_client_number or "").strip()
                    if existing_parent
                    else ""
                ) or f"client #{global_row.client_id}"
                uploaded_parent = (
                    self.db.query(PayflowClientModel)
                    .filter(PayflowClientModel.id == client_id)
                    .first()
                )
                uploaded_master = (
                    (uploaded_parent.crm_client_number or "").strip()
                    if uploaded_parent
                    else ""
                ) or f"client #{client_id}"
                raise ValidationAppError(
                    f'Sub-client "{number}" already belongs to master '
                    f'"{existing_master}". It cannot be moved under '
                    f'"{uploaded_master}".'
                )
            if global_row is not None:
                return global_row
        if code_norm:
            by_code = (
                self.db.query(PayflowPortfolioModel)
                .filter(
                    PayflowPortfolioModel.client_id == client_id,
                    PayflowPortfolioModel.code == code_norm,
                )
                .first()
            )
            if by_code is None:
                return None
            existing_number = (by_code.crm_client_number or "").strip()
            if number and existing_number and existing_number.lower() != number.lower():
                raise ValidationAppError(
                    f'Sub-client number cannot be changed for an existing '
                    f'client_code "{code_norm}". Keep the original number '
                    f'"{existing_number}", or use a new number and a new code '
                    "to create a new sub-client."
                )
            return by_code
        return None

    def _find_portfolio(
        self,
        client_id: int,
        *,
        crm_number: str | None,
        code: str | None,
    ) -> PayflowPortfolioModel | None:
        if crm_number:
            row = (
                self.db.query(PayflowPortfolioModel)
                .filter(
                    PayflowPortfolioModel.client_id == client_id,
                    PayflowPortfolioModel.crm_client_number == crm_number,
                )
                .first()
            )
            if row:
                return row
        if code:
            return (
                self.db.query(PayflowPortfolioModel)
                .filter(
                    PayflowPortfolioModel.client_id == client_id,
                    PayflowPortfolioModel.code == code,
                )
                .first()
            )
        return None

    @staticmethod
    def _normalize_data_source(value: Any, *, allow_none: bool = False) -> str | None:
        if value is None or value == "":
            if allow_none:
                return None
            raise ValidationAppError("Data source is required")
        ds = str(value).strip().lower()
        if ds in ("crm", PayflowDataSourceType.CRM.value):
            return PayflowDataSourceType.CRM.value
        if ds in ("file", "daily_file", "file_only", PayflowDataSourceType.FILE.value):
            return PayflowDataSourceType.FILE.value
        raise ValidationAppError("Data source must be CRM or File")

    def _set_data_source(self, client: PayflowClientModel, ds: str | None) -> None:
        prev = client.data_source_type
        client.data_source_type = ds
        if ds is None:
            client.connection_status = PayflowConnectionStatus.NOT_CONNECTED.value
        elif ds == PayflowDataSourceType.FILE.value:
            client.connection_status = PayflowConnectionStatus.CONNECTED.value
        elif ds == PayflowDataSourceType.CRM.value and prev != PayflowDataSourceType.CRM.value:
            client.connection_status = PayflowConnectionStatus.NOT_CONNECTED.value

    def _apply_primary_fields(
        self, client: PayflowClientModel, payload: dict[str, Any], *, create: bool
    ) -> None:
        for field in _PRIMARY_CONTACT_FIELDS:
            if field not in payload:
                continue
            val = payload.get(field)
            if val is None and not create:
                continue
            setattr(client, field, (str(val).strip() if val is not None else None) or None)
        if payload.get("integration_ref") is not None:
            client.integration_ref = str(payload["integration_ref"]).strip() or None
        elif create and payload.get("crm_client_number"):
            client.integration_ref = str(payload["crm_client_number"]).strip()
        if payload.get("crm_system_name") is not None:
            client.crm_system_name = str(payload["crm_system_name"]).strip() or None

    def _seed_default_mappings(self, client: PayflowClientModel) -> None:
        for i, field in enumerate(CRM_INBOUND_FIELDS):
            status, payflow_field = self._default_mapping_state(field)
            self.db.add(
                PayflowClientFieldMappingModel(
                    client_id=client.id,
                    source_field=field["source_field"],
                    payflow_field=payflow_field,
                    sample_value=field["sample_value"] or None,
                    status=status,
                    is_required=field["required"],
                    sort_order=i,
                )
            )
        self.db.flush()

    @staticmethod
    def _default_mapping_state(field: dict) -> tuple[str, str]:
        """Return (status, payflow_field) for a catalog field."""
        available = (field.get("available") or "").strip()
        payflow = field["payflow_field"]
        if available == "PayFlow-built":
            return PayflowMappingStatus.UNMAPPED.value, "— Not mapped —"
        # Open decisions / missing / partial columns → Needs Attention
        if available in ("Ambiguous", "No", "Partial"):
            if field["required"] or available == "Ambiguous":
                return PayflowMappingStatus.NEEDS_ATTENTION.value, payflow
            return PayflowMappingStatus.UNMAPPED.value, "— Not mapped —"
        if field["required"] or available in ("Yes", "Derived"):
            return PayflowMappingStatus.MAPPED.value, payflow
        return PayflowMappingStatus.UNMAPPED.value, "— Not mapped —"

    def _sync_catalog_mappings(self, client: PayflowClientModel) -> bool:
        """Align stored mappings with CRM catalog v0.3. Returns True if changed."""
        existing = {
            m.source_field: m
            for m in self.db.query(PayflowClientFieldMappingModel)
            .filter(PayflowClientFieldMappingModel.client_id == client.id)
            .all()
        }
        catalog_sources = {f["source_field"] for f in CRM_INBOUND_FIELDS}
        changed = False

        # Drop obsolete v0.1 rows not in v0.3 catalog.
        for source, row in list(existing.items()):
            if source not in catalog_sources:
                self.db.delete(row)
                existing.pop(source, None)
                changed = True

        for i, field in enumerate(CRM_INBOUND_FIELDS):
            source = field["source_field"]
            status, payflow_field = self._default_mapping_state(field)
            row = existing.get(source)
            if row is None:
                self.db.add(
                    PayflowClientFieldMappingModel(
                        client_id=client.id,
                        source_field=source,
                        payflow_field=payflow_field,
                        sample_value=field["sample_value"] or None,
                        status=status,
                        is_required=field["required"],
                        sort_order=i,
                    )
                )
                changed = True
                continue
            # Refresh catalog metadata; preserve admin Mapped/Validated choices when still valid.
            if row.is_required != field["required"]:
                row.is_required = field["required"]
                changed = True
            if row.sort_order != i:
                row.sort_order = i
                changed = True
            sample = field["sample_value"] or None
            if row.sample_value != sample:
                row.sample_value = sample
                changed = True
            # If still pointing at a removed PayFlow label, reset to catalog default.
            pf = (row.payflow_field or "").strip()
            valid_targets = set(PAYFLOW_TARGET_FIELDS)
            if pf and pf not in valid_targets:
                row.payflow_field = payflow_field
                row.status = status
                changed = True
            elif row.status == PayflowMappingStatus.UNMAPPED.value and field["required"]:
                row.payflow_field = payflow_field
                row.status = status
                changed = True

        if changed:
            self.db.flush()
        return changed

    def _replace_mappings(self, client: PayflowClientModel, mappings: list) -> None:
        if not isinstance(mappings, list):
            raise ValidationAppError("mappings must be a list")
        existing = {
            m.source_field: m
            for m in self.db.query(PayflowClientFieldMappingModel)
            .filter(PayflowClientFieldMappingModel.client_id == client.id)
            .all()
        }
        for item in mappings:
            source = str(item.get("source_field") or "").strip()
            if not source or source not in existing:
                raise ValidationAppError(f"Unknown source field: {source or '(empty)'}")
            row = existing[source]
            if "payflow_field" in item and item["payflow_field"] is not None:
                row.payflow_field = str(item["payflow_field"]).strip()
            if "sample_value" in item and item["sample_value"] is not None:
                row.sample_value = str(item["sample_value"]).strip() or None
            if "status" in item and item["status"] is not None:
                st = str(item["status"]).strip().lower()
                if st not in _MAPPING_FROM_LABEL:
                    raise ValidationAppError(f"Invalid mapping status: {item['status']}")
                row.status = _MAPPING_FROM_LABEL[st]
            else:
                # Derive status
                pf = (row.payflow_field or "").strip()
                if not pf or pf == "— Not mapped —":
                    row.status = PayflowMappingStatus.UNMAPPED.value
                else:
                    row.status = PayflowMappingStatus.MAPPED.value

    def _sync_supervisors(self, client: PayflowClientModel, user_ids: list) -> list:
        try:
            wanted = {UUID(str(u)) for u in user_ids}
        except Exception as exc:
            raise ValidationAppError("Invalid supervisor_user_ids") from exc

        newly_assigned: list = []

        # Current assignments for this client
        current = (
            self.db.query(PayflowUserClientAssignmentModel)
            .join(PayflowUserMembershipModel)
            .options(joinedload(PayflowUserClientAssignmentModel.membership).joinedload(
                PayflowUserMembershipModel.role
            ))
            .filter(PayflowUserClientAssignmentModel.client_id == client.id)
            .all()
        )
        current_by_user: dict[UUID, PayflowUserClientAssignmentModel] = {}
        for a in current:
            if a.membership and a.membership.user_id:
                current_by_user[a.membership.user_id] = a

        # Remove client-scoped users no longer wanted (never touch platform-wide admins).
        for uid, assignment in list(current_by_user.items()):
            role = assignment.membership.role if assignment.membership else None
            scope = role.scope if role else None
            if uid not in wanted and scope == PayflowRoleScope.CLIENT_SCOPED.value:
                self.db.delete(assignment)

        # Add missing — any client-scoped role (Supervisor or custom).
        for uid in wanted:
            if uid in current_by_user:
                continue
            membership = (
                self.db.query(PayflowUserMembershipModel)
                .options(joinedload(PayflowUserMembershipModel.role))
                .filter(PayflowUserMembershipModel.user_id == uid)
                .first()
            )
            if not membership:
                raise ValidationAppError(f"User {uid} has no PayFlow membership")
            if not membership.role or membership.role.scope != PayflowRoleScope.CLIENT_SCOPED.value:
                raise ValidationAppError(
                    "Only client-scoped users can be assigned to a client"
                )
            user = self.db.query(UserModel).filter(UserModel.id == uid).first()
            if not user or user.status == "disabled":
                raise ValidationAppError("User is not available for client assignment")
            assignment = PayflowUserClientAssignmentModel(
                membership_id=membership.id,
                client_id=client.id,
            )
            self.db.add(assignment)
            self.db.flush()
            # Permissions come live from the role — no per-assignment snapshot.
            newly_assigned.append(user)
        self.db.flush()
        return newly_assigned

    def _supervisors(self, client: PayflowClientModel) -> list[dict]:
        result = []
        for a in client.assignments or []:
            m = a.membership
            if not m or not m.user:
                continue
            if m.role and m.role.scope == PayflowRoleScope.PLATFORM_WIDE.value:
                continue
            # Always show current role permissions (Edit access updates apply immediately).
            perm_names: list[str] = []
            if m.role:
                role_links = (
                    self.db.query(PayflowRolePermissionModel)
                    .options(joinedload(PayflowRolePermissionModel.permission))
                    .filter(PayflowRolePermissionModel.payflow_role_id == m.role.id)
                    .all()
                )
                for link in role_links:
                    if link.permission and link.permission.name:
                        perm_names.append(link.permission.name)
            result.append(
                {
                    "user_id": str(m.user_id),
                    "full_name": m.user.full_name,
                    "email": m.user.email,
                    "status": m.user.status,
                    "short_name": (m.user.full_name or "").split()[0] if m.user.full_name else "",
                    "role_name": m.role.name if m.role else None,
                    "permission_names": sorted(set(perm_names)),
                }
            )
        return result

    def _mapping_summary(self, client: PayflowClientModel) -> dict:
        mappings = list(client.field_mappings or [])
        mapped = sum(
            1
            for m in mappings
            if m.status
            in (PayflowMappingStatus.MAPPED.value, PayflowMappingStatus.VALIDATED.value)
        )
        attention = sum(
            1 for m in mappings if m.status == PayflowMappingStatus.NEEDS_ATTENTION.value
        )
        unmapped = sum(
            1 for m in mappings if m.status == PayflowMappingStatus.UNMAPPED.value
        )
        required_missing = [
            m.source_field
            for m in mappings
            if m.is_required
            and m.status
            not in (PayflowMappingStatus.MAPPED.value, PayflowMappingStatus.VALIDATED.value)
        ]
        return {
            "mapped": mapped,
            "attention": attention,
            "unmapped": unmapped,
            "total": len(mappings),
            "required_missing": required_missing,
        }

    def _profile_issues(self, client: PayflowClientModel) -> list[str]:
        """General / Client Profile field issues (same rules as FE General tab)."""
        issues: list[str] = []
        if not (client.name or "").strip():
            issues.append("Client name is required")
        if not (client.code or "").strip():
            issues.append("Client code is required")
        if not (client.client_type or "").strip():
            issues.append("Client type is required")
        if not (client.business_domain or "").strip():
            issues.append("Business use case is required")
        if not (client.contact_name or "").strip():
            issues.append("Primary contact name is required")
        email = (client.contact_email or "").strip()
        if not email:
            issues.append("Primary contact email is required")
        elif not _EMAIL_RE.match(email):
            issues.append("Enter a valid primary contact email")
        if not (client.contact_phone or "").strip():
            issues.append("Primary contact phone is required")
        if not (client.country or "").strip():
            issues.append("Country is required")
        if not (client.correspondence_language or "").strip():
            issues.append("Language is required")
        if not (client.currency_code or "").strip():
            issues.append("Currency is required")
        postal = (client.postal_code or "").strip()
        if postal and not _is_valid_postal_code(postal):
            issues.append(
                "Postal code must match format A1A 1A1 "
                "(letter-number-letter space number-letter-number)"
            )
        return issues

    def _onboarding_progress(self, client: PayflowClientModel) -> dict:
        supervisors = self._supervisors(client)
        portfolio_count = len(client.portfolios or [])

        profile_ok = len(self._profile_issues(client)) == 0
        # Clients use daily file intake; CRM field mapping is global (System Mapping).
        is_file = client.data_source_type == PayflowDataSourceType.FILE.value
        is_crm = client.data_source_type == PayflowDataSourceType.CRM.value
        source_ok = is_file or (
            is_crm and client.connection_status == PayflowConnectionStatus.CONNECTED.value
        )
        branding_ok = len(self._branding_issues(client)) == 0
        ai_ok = bool(client.ai_mode)
        supervisors_ok = len(supervisors) > 0
        activation_ok = client.status == PayflowClientStatus.ACTIVE.value
        # Data Mapping is system-fixed (Administration → System Mapping) — not a client journey step.

        def step(ok: bool, started: bool = True) -> str:
            if ok:
                return PayflowOnboardingStepStatus.COMPLETE.value
            if not started:
                return PayflowOnboardingStepStatus.PENDING.value
            return PayflowOnboardingStepStatus.INCOMPLETE.value

        steps = [
            {
                "key": "profile",
                "label": "Client Profile",
                "status": step(profile_ok, True),
            },
            {
                "key": "portfolios",
                "label": "Sub-Clients / Portfolios",
                "status": (
                    PayflowOnboardingStepStatus.COMPLETE.value
                    if portfolio_count > 0
                    else PayflowOnboardingStepStatus.PENDING.value
                ),
                "informational": True,
            },
            {
                "key": "data_source",
                "label": "Data Source",
                "status": step(source_ok, client.data_source_type is not None),
            },
            {
                "key": "branding",
                "label": "Branding & Channels",
                "status": step(branding_ok, True),
            },
            {
                "key": "ai_governance",
                "label": "AI & Governance",
                "status": step(ai_ok, True),
            },
            {
                "key": "supervisors",
                "label": "Supervisor Assignment",
                "status": step(supervisors_ok, True),
            },
            {
                "key": "activation",
                "label": "Review & Activation",
                "status": (
                    PayflowOnboardingStepStatus.COMPLETE.value
                    if activation_ok
                    else (
                        PayflowOnboardingStepStatus.BLOCKED.value
                        if self._activation_blockers(client)
                        else PayflowOnboardingStepStatus.PENDING.value
                    )
                ),
            },
        ]
        required_complete = sum(
            1
            for s in steps
            if not s.get("informational")
            and s["status"] == PayflowOnboardingStepStatus.COMPLETE.value
        )
        required_total = sum(1 for s in steps if not s.get("informational"))
        return {
            "steps": steps,
            "completed_required": required_complete,
            "total_required": required_total,
            "percent": int(round((required_complete / required_total) * 100)) if required_total else 0,
            "eligible_for_activation": len(self._activation_blockers(client)) == 0,
        }

    def _branding_issues(self, client: PayflowClientModel) -> list[str]:
        """Required branding/channel fields — aligned with FE Branding & Channels save rules."""
        issues: list[str] = []
        if not client.channel_email and not client.channel_sms:
            issues.append("At least one communication channel must be enabled")
        brand = (client.brand_name or "").strip()
        if not brand:
            issues.append("Display / brand name is required (set in Branding & Channels)")
        if not (client.sender_name or "").strip():
            issues.append("Sender name is required (set in Branding & Channels)")
        is_first = client.client_type == PayflowClientType.FIRST_PARTY.value
        if is_first:
            if client.channel_email and not (client.email_from or "").strip():
                issues.append("Email from address is required when Email is enabled")
            if client.channel_sms and not (client.sms_sender_id or "").strip():
                issues.append("SMS sender ID is required when SMS is enabled")
        return issues

    def _setup_issue_messages(
        self, client: PayflowClientModel, *, include_informational: bool = True
    ) -> list[str]:
        """Concrete remaining setup issues across all config tabs (for grid / header)."""
        issues: list[str] = []
        progress = self._onboarding_progress(client)
        by_key = {s["key"]: s for s in progress["steps"]}

        for msg in self._profile_issues(client):
            issues.append(f"Client Profile: {msg}")

        if include_informational:
            port_status = by_key.get("portfolios", {}).get("status")
            if port_status != PayflowOnboardingStepStatus.COMPLETE.value:
                issues.append(
                    "Sub-Clients / Portfolios: add at least one sub-client (optional but recommended)"
                )

        if by_key.get("data_source", {}).get("status") != PayflowOnboardingStepStatus.COMPLETE.value:
            issues.append("Data Source: select and connect a primary data source")

        for msg in self._branding_issues(client):
            issues.append(f"Branding & Channels: {msg}")

        if by_key.get("ai_governance", {}).get("status") != PayflowOnboardingStepStatus.COMPLETE.value:
            issues.append("AI & Governance: select an AI mode")

        if by_key.get("supervisors", {}).get("status") != PayflowOnboardingStepStatus.COMPLETE.value:
            issues.append("Supervisor Assignment: assign at least one supervisor")

        act = by_key.get("activation", {})
        if act.get("status") != PayflowOnboardingStepStatus.COMPLETE.value:
            if act.get("status") == PayflowOnboardingStepStatus.BLOCKED.value:
                for b in self._activation_blockers(client):
                    label = f"Review & Activation: {b}"
                    if label not in issues and not any(
                        b.lower() in (i or "").lower() for i in issues
                    ):
                        issues.append(label)
            else:
                issues.append(
                    "Review & Activation: finish required setup, then activate the client"
                )

        # De-dupe while preserving order
        seen: set[str] = set()
        unique: list[str] = []
        for item in issues:
            key = item.lower()
            if key in seen:
                continue
            seen.add(key)
            unique.append(item)
        return unique

    def _simulate_setup_issues_after_file_import(
        self,
        payload: dict[str, Any],
        *,
        has_subs_in_file: bool = False,
    ) -> list[str]:
        """Expected remaining issues for a master created/updated from CRM file (defaults)."""
        client_type = _parse_client_type_soft(payload.get("client_type"))
        brand = (
            (payload.get("brand_name") or "").strip()
            or (payload.get("short_name") or "").strip()
            or (payload.get("name") or "").strip()
        )
        # File import always seeds empty sender_name and no supervisors.
        stub = PayflowClientModel(
            name=(payload.get("name") or "").strip() or "Client",
            code=_normalize_code(payload.get("code") or "") or "TMP",
            client_type=client_type,
            business_domain=_parse_domain(payload.get("business_domain")),
            ai_mode=_parse_ai_mode(payload.get("ai_mode")),
            data_source_type=PayflowDataSourceType.CRM.value,
            connection_status=PayflowConnectionStatus.CONNECTED.value,
            contact_name=(payload.get("contact_name") or "").strip() or None,
            contact_email=(payload.get("contact_email") or "").strip() or None,
            contact_phone=(payload.get("contact_phone") or "").strip() or None,
            country=(payload.get("country") or "").strip() or None,
            correspondence_language=(
                payload.get("correspondence_language") or ""
            ).strip()
            or None,
            currency_code=(payload.get("currency_code") or "").strip() or None,
            postal_code=(payload.get("postal_code") or "").strip() or None,
            brand_name=brand,
            sender_name="",
            email_from="collections@payflow.io",
            sms_sender_id="PAYFLOW",
            channel_email=True,
            channel_sms=True,
            status=PayflowClientStatus.DRAFT.value,
            portfolios=[],
        )
        # Avoid DB hits in _supervisors / portfolio count — stub has no id/assignments
        issues: list[str] = []
        for msg in self._profile_issues(stub):
            issues.append(f"Client Profile: {msg}")
        for msg in self._branding_issues(stub):
            issues.append(f"Branding & Channels: {msg}")
        issues.append("Supervisor Assignment: assign at least one supervisor")
        if not has_subs_in_file:
            issues.append(
                "Sub-Clients / Portfolios: add at least one sub-client (optional but recommended)"
            )
        issues.append(
            "Review & Activation: finish required setup, then activate the client"
        )
        return issues

    def _require_draft_for_branding(self, client: PayflowClientModel) -> None:
        if client.status != PayflowClientStatus.DRAFT.value:
            raise ValidationAppError(
                "Branding & channels can only be updated while the client is in Draft status"
            )

    def _activation_blockers(self, client: PayflowClientModel) -> list[str]:
        blockers: list[str] = []
        blockers.extend(self._profile_issues(client))
        # Phase-1 client intake is daily file; CRM→PayFlow mapping is system-wide.
        if client.data_source_type not in (
            PayflowDataSourceType.FILE.value,
            PayflowDataSourceType.CRM.value,
        ):
            blockers.append("A primary data source has not been selected")
        blockers.extend(self._branding_issues(client))
        if not self._supervisors(client):
            blockers.append("Assign at least one supervisor")
        return blockers

    def _list_item(self, client: PayflowClientModel) -> dict:
        # Light load supervisors without full joinedload on list — query assignments
        supervisors = []
        assignments = (
            self.db.query(PayflowUserClientAssignmentModel)
            .options(
                joinedload(PayflowUserClientAssignmentModel.membership)
                .joinedload(PayflowUserMembershipModel.user),
                joinedload(PayflowUserClientAssignmentModel.membership)
                .joinedload(PayflowUserMembershipModel.role),
            )
            .filter(PayflowUserClientAssignmentModel.client_id == client.id)
            .all()
        )
        for a in assignments:
            m = a.membership
            if not m or not m.user:
                continue
            if m.role and m.role.code == PayflowRoleCode.OPERATIONS_ADMIN.value:
                continue
            supervisors.append(
                {
                    "user_id": str(m.user_id),
                    "full_name": m.user.full_name,
                    "short_name": (m.user.full_name or "").split()[0] if m.user.full_name else "",
                }
            )

        progress = self._onboarding_progress(client)
        # All journey steps still open (Data Mapping is not in steps).
        incomplete = [
            s["label"]
            for s in progress["steps"]
            if s.get("status") != PayflowOnboardingStepStatus.COMPLETE.value
        ]
        setup_issues = self._setup_issue_messages(client)
        return {
            "id": client.id,
            "code": client.code,
            "name": client.name,
            "category": client.category,
            "industry": client.category,
            "status": client.status,
            "status_label": _STATUS_LABELS.get(client.status, client.status),
            "client_type": client.client_type,
            "client_type_label": _CLIENT_TYPE_LABELS.get(client.client_type or "", client.client_type),
            "business_domain": client.business_domain,
            "business_domain_label": _DOMAIN_LABELS.get(
                client.business_domain or "", client.business_domain
            ),
            "ai_mode": client.ai_mode,
            "ai_mode_label": _AI_MODE_LABELS.get(client.ai_mode or "", client.ai_mode),
            "data_source_type": client.data_source_type,
            "connection_status": client.connection_status,
            "connection_status_label": _CONNECTION_LABELS.get(
                client.connection_status or "", client.connection_status
            ),
            "supervisors": supervisors,
            "setup_incomplete": incomplete,
            "setup_issues": setup_issues,
            "setup_steps_remaining": max(len(incomplete), len(setup_issues)),
            "onboarding": progress,
            "added_through": client.added_through
            or PayflowClientAddedThrough.ADD_CLIENT.value,
            "added_through_label": _ADDED_THROUGH_LABELS.get(
                client.added_through or PayflowClientAddedThrough.ADD_CLIENT.value,
                "Created",
            ),
            "created_at": client.created_at,
            "updated_at": client.updated_at or client.created_at,
        }

    def _detail(self, client: PayflowClientModel) -> dict:
        # Ensure relationships loaded
        client = self._get_or_404(client.id)
        summary = self._mapping_summary(client)
        supervisors = self._supervisors(client)
        progress = self._onboarding_progress(client)
        blockers = self._activation_blockers(client)

        mappings = sorted(client.field_mappings or [], key=lambda m: m.sort_order)
        portfolios = sorted(client.portfolios or [], key=lambda p: p.name.lower())

        return {
            **self._list_item(client),
            "crm_system_name": client.crm_system_name,
            "integration_ref": client.integration_ref,
            "crm_client_number": client.crm_client_number,
            "contact_name": client.contact_name,
            "contact_title": client.contact_title,
            "contact_email": client.contact_email,
            "contact_phone": client.contact_phone,
            "address_line1": client.address_line1,
            "address_line2": client.address_line2,
            "city": client.city,
            "province_state": client.province_state,
            "country": client.country,
            "postal_code": client.postal_code,
            "correspondence_language": client.correspondence_language,
            "currency_code": client.currency_code,
            "crm_status": client.crm_status,
            "environment": client.environment,
            "sync_frequency": client.sync_frequency,
            "brand_name": client.brand_name or "",
            "logo_url": client.logo_url,
            "sender_name": client.sender_name or "",
            "email_from": client.email_from or "",
            "sms_sender_id": client.sms_sender_id or "",
            "channels": {
                "email": bool(client.channel_email),
                "sms": bool(client.channel_sms),
                "whatsapp": bool(client.channel_whatsapp),
            },
            "governance_rules": list(client.governance_rules or []),
            "mappings": [
                {
                    "id": m.id,
                    "source_field": m.source_field,
                    "payflow_field": m.payflow_field,
                    "sample_value": m.sample_value,
                    "status": m.status,
                    "status_label": _MAPPING_LABELS.get(m.status, m.status),
                    "is_required": m.is_required,
                    "sort_order": m.sort_order,
                }
                for m in mappings
            ],
            "mapping_summary": summary,
            "portfolios": [self._portfolio_item(p) for p in portfolios],
            "portfolio_count": len(portfolios),
            "supervisor_user_ids": [s["user_id"] for s in supervisors],
            "supervisors": supervisors,
            "onboarding": progress,
            "activation_blockers": blockers,
        }

    def _portfolio_item(self, p: PayflowPortfolioModel) -> dict:
        accounts = list(p.accounts or [])
        account_count = len(accounts)
        case_count = account_count  # one collection case per account in Phase 1
        outstanding = float(sum(float(a.outstanding_balance or 0) for a in accounts))
        active_strategy = None
        for s in p.strategies or []:
            if (s.status or "").lower() in ("active", "approved"):
                active_strategy = s
                break
        if active_strategy is None and (p.strategies or []):
            # Prefer most recently updated strategy for display if none active
            active_strategy = max(
                p.strategies,
                key=lambda s: s.updated_at or s.created_at,
            )
            if (active_strategy.status or "").lower() not in ("active", "approved"):
                active_strategy = None

        return {
            "id": p.id,
            "client_id": p.client_id,
            "name": p.name,
            "code": p.code,
            "status": p.status,
            "status_label": {
                PayflowPortfolioStatus.ONBOARDING.value: "Onboarding",
                PayflowPortfolioStatus.ACTIVE.value: "Active",
                PayflowPortfolioStatus.PAUSED.value: "Paused",
            }.get(p.status, p.status),
            "description": p.description,
            "crm_client_number": p.crm_client_number,
            "account_count": account_count,
            "case_count": case_count,
            "outstanding": outstanding,
            "active_strategy_id": active_strategy.id if active_strategy else None,
            "active_strategy_name": active_strategy.name if active_strategy else None,
            "last_file_received": "No file received yet",
            "created_at": p.created_at,
            "updated_at": p.updated_at,
        }

    def get_portfolio(self, client_id: int, portfolio_id: int) -> dict:
        client = self._get_or_404(client_id)
        portfolio = (
            self.db.query(PayflowPortfolioModel)
            .options(
                joinedload(PayflowPortfolioModel.accounts),
                joinedload(PayflowPortfolioModel.strategies),
            )
            .filter(
                PayflowPortfolioModel.id == portfolio_id,
                PayflowPortfolioModel.client_id == client_id,
            )
            .first()
        )
        if not portfolio:
            raise NotFoundError("Portfolio not found")
        strategies = sorted(
            portfolio.strategies or [],
            key=lambda s: s.updated_at or s.created_at,
            reverse=True,
        )
        return {
            **self._portfolio_item(portfolio),
            "client_name": client.name,
            "client_code": client.code,
            "strategies": [
                {
                    "id": s.id,
                    "name": s.name,
                    "code": s.code,
                    "status": s.status,
                    "origin": s.origin,
                    "source": getattr(s, "source", None) or "Human Created",
                    "version": s.version,
                    "updated_at": s.updated_at or s.created_at,
                }
                for s in strategies
            ],
        }
